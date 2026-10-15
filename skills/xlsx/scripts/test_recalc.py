"""
Unit and regression tests for skills/xlsx/scripts/recalc.py
Validates external_links_at_risk, defined name scopes, Unicode aliases,
sheet qualifier handling, and cycle protection.
"""

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from openpyxl.packaging.relationship import Relationship
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.external_link.external import (
    ExternalBook,
    ExternalLink,
    ExternalSheetNames,
)

# Import the module under test
import recalc


def attach_external_link(wb, target="file:///missing-source.xlsx", sheet_name="Sheet1"):
    """Attaches a genuine external link part to the workbook."""
    ext = ExternalLink(
        externalBook=ExternalBook(
            sheetNames=ExternalSheetNames(sheetName=[sheet_name]), id="rId1"
        )
    )
    ext.file_link = Relationship(
        Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLinkPath",
        Target=target,
        TargetMode="External",
        Id="rId1",
    )
    wb._external_links.append(ext)


class TestExternalLinksAtRisk(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.tmp_path = Path(self.temp_dir.name)

    def tearDown(self):
        self.temp_dir.cleanup()

    def _save_and_check(self, wb, filename="test.xlsx"):
        path = self.tmp_path / filename
        wb.save(path)
        wb.close()
        return recalc.external_links_at_risk(path)

    # -------------------------------------------------------------
    # Phase 2: Fix #1 - Unicode defined-name aliases
    # -------------------------------------------------------------
    def test_unicode_defined_name_alias(self):
        """Alias chain through Unicode name: AliasU -> Übersicht -> external."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Main"
        attach_external_link(wb)
        wb.defined_names.add(DefinedName("Übersicht", attr_text="'[1]Sheet1'!$C$1"))
        wb.defined_names.add(DefinedName("AliasU", attr_text="Übersicht"))
        ws["A1"] = "=AliasU"
        result = self._save_and_check(wb, "unicode_alias.xlsx")
        self.assertEqual(result, ["Main!A1"])

    def test_unicode_multi_hop_alias(self):
        """Multi-hop alias chain with Unicode names: TopU -> MidU -> Übersicht -> external."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Main"
        attach_external_link(wb)
        wb.defined_names.add(DefinedName("Übersicht", attr_text="'[1]Sheet1'!$C$1"))
        wb.defined_names.add(DefinedName("MidU", attr_text="Übersicht"))
        wb.defined_names.add(DefinedName("TopU", attr_text="MidU"))
        ws["A1"] = "=TopU"
        result = self._save_and_check(wb, "unicode_multi_hop.xlsx")
        self.assertEqual(result, ["Main!A1"])

    # -------------------------------------------------------------
    # Phase 3: Fix #2 - Prevent false positives from sheet names
    # -------------------------------------------------------------
    def test_sheet_name_containing_external_token(self):
        """Sheet name contains token matching external defined name (no false positive)."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Main"
        wb.create_sheet("EXT Sheet")
        attach_external_link(wb)
        # EXT points externally
        wb.defined_names.add(DefinedName("EXT", attr_text="'[1]Sheet1'!$A$1"))
        # Rate points to local sheet 'EXT Sheet'
        wb.defined_names.add(DefinedName("Rate", attr_text="'EXT Sheet'!$A$1"))
        ws["A1"] = "=Rate"
        ws["A2"] = "='EXT Sheet'!A1"
        ws["A3"] = "=EXT"  # genuine control
        result = self._save_and_check(wb, "sheet_name_token.xlsx")
        self.assertNotIn("Main!A1", result)
        self.assertNotIn("Main!A2", result)
        self.assertIn("Main!A3", result)

    def test_quoted_sheet_with_spaces_and_special_chars(self):
        """Quoted sheet names with spaces, special characters, or dots."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Main"
        wb.create_sheet("2026 Report")
        attach_external_link(wb)
        wb.defined_names.add(DefinedName("Report", attr_text="'2026 Report'!$B$5"))
        ws["A1"] = "=Report"
        ws["A2"] = "='2026 Report'!$B$5"
        result = self._save_and_check(wb, "quoted_sheet.xlsx")
        self.assertEqual(result, [])

    def test_string_literal_not_flagged(self):
        """String literals matching defined names in formulas must not be flagged."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Main"
        attach_external_link(wb)
        wb.defined_names.add(DefinedName("EXT", attr_text="'[1]Sheet1'!$A$1"))
        ws["A1"] = '="EXT"'
        ws["A2"] = '="EXT" & " Data"'
        result = self._save_and_check(wb, "string_literal.xlsx")
        self.assertEqual(result, [])

    def test_internal_only_alias(self):
        """Ordinary internal aliases must remain unflagged."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Main"
        attach_external_link(wb)
        wb.defined_names.add(DefinedName("Ordinary", attr_text="Main!$A$2"))
        ws["A1"] = "=Ordinary"
        ws["A2"] = 100
        result = self._save_and_check(wb, "internal_alias.xlsx")
        self.assertEqual(result, [])

    # -------------------------------------------------------------
    # Phase 4: Fix #3 - Cross-worksheet scoped defined names
    # -------------------------------------------------------------
    def test_cross_sheet_scoped_name(self):
        """Cross-sheet qualified reference: ='Source Data'!ExternalAmount."""
        wb = Workbook()
        main = wb.active
        main.title = "Main"
        source = wb.create_sheet("Source Data")
        attach_external_link(wb)
        source.defined_names.add(
            DefinedName("ExternalAmount", attr_text="'[1]Sheet1'!$A$1", localSheetId=1)
        )
        main["A1"] = "='Source Data'!ExternalAmount"
        result = self._save_and_check(wb, "cross_sheet.xlsx")
        self.assertEqual(result, ["Main!A1"])

    def test_alias_pointing_to_qualified_name(self):
        """Workbook-level alias pointing to cross-sheet qualified name."""
        wb = Workbook()
        main = wb.active
        main.title = "Main"
        source = wb.create_sheet("Source Data")
        attach_external_link(wb)
        source.defined_names.add(
            DefinedName("ExternalAmount", attr_text="'[1]Sheet1'!$A$1", localSheetId=1)
        )
        wb.defined_names.add(
            DefinedName("GlobalAlias", attr_text="'Source Data'!ExternalAmount")
        )
        main["A1"] = "=GlobalAlias"
        result = self._save_and_check(wb, "alias_qualified.xlsx")
        self.assertEqual(result, ["Main!A1"])

    def test_current_sheet_shadowing(self):
        """Worksheet-level name shadows workbook/other sheet for unqualified references."""
        wb = Workbook()
        main = wb.active
        main.title = "Main"
        source = wb.create_sheet("Source Data")
        attach_external_link(wb)
        # Main has local ExternalAmount = 42 (internal)
        main.defined_names.add(
            DefinedName("ExternalAmount", attr_text="42", localSheetId=0)
        )
        # Source Data has local ExternalAmount pointing to external
        source.defined_names.add(
            DefinedName("ExternalAmount", attr_text="'[1]Sheet1'!$A$1", localSheetId=1)
        )
        main["A1"] = "=ExternalAmount"              # Local to Main -> 42 -> safe
        main["A2"] = "='Source Data'!ExternalAmount"  # Explicitly qualified -> at risk
        source["A1"] = "=ExternalAmount"              # Local to Source Data -> at risk
        result = self._save_and_check(wb, "shadowing.xlsx")
        self.assertNotIn("Main!A1", result)
        self.assertIn("Main!A2", result)
        self.assertIn("Source Data!A1", result)

    def test_workbook_scope_fallback_and_local_precedence(self):
        """Unqualified name falls back to workbook scope if not defined on sheet."""
        wb = Workbook()
        main = wb.active
        main.title = "Main"
        source = wb.create_sheet("Source Data")
        attach_external_link(wb)
        # Global TaxRate is safe
        wb.defined_names.add(DefinedName("TaxRate", attr_text="0.1"))
        # Source Data has local TaxRate pointing to external
        source.defined_names.add(
            DefinedName("TaxRate", attr_text="'[1]Sheet1'!$B$2", localSheetId=1)
        )
        main["A1"] = "=TaxRate"   # Resolves to global 0.1 -> safe
        source["A1"] = "=TaxRate" # Resolves to local external -> at risk
        result = self._save_and_check(wb, "scope_precedence.xlsx")
        self.assertNotIn("Main!A1", result)
        self.assertIn("Source Data!A1", result)

    # -------------------------------------------------------------
    # Phase 5: Preserve existing behavior
    # -------------------------------------------------------------
    def test_direct_external_reference(self):
        """Direct external reference formula in cell."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Main"
        attach_external_link(wb)
        ws["A1"] = "='[1]Sheet1'!$A$1"
        result = self._save_and_check(wb, "direct_ext.xlsx")
        self.assertEqual(result, ["Main!A1"])

    def test_case_insensitive_name_reference(self):
        """Defined names match case-insensitively."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Main"
        attach_external_link(wb)
        wb.defined_names.add(DefinedName("CaseRef", attr_text="'[1]Sheet1'!$A$1"))
        ws["A1"] = "=caseref"
        result = self._save_and_check(wb, "case_insensitive.xlsx")
        self.assertEqual(result, ["Main!A1"])

    def test_cyclic_alias_chain_safe(self):
        """Cyclic alias chain terminates safely without recursion error or hang."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Main"
        attach_external_link(wb)
        wb.defined_names.add(DefinedName("CycA", attr_text="CycB"))
        wb.defined_names.add(DefinedName("CycB", attr_text="CycA"))
        ws["A1"] = "=CycA"
        ws["A2"] = "=1+1"
        result = self._save_and_check(wb, "cyclic.xlsx")
        self.assertEqual(result, [])

    def test_ordinary_math_formula(self):
        """Ordinary formula =1+1 is not flagged."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Main"
        attach_external_link(wb)
        ws["A1"] = "=1+1"
        result = self._save_and_check(wb, "math.xlsx")
        self.assertEqual(result, [])

    def test_no_external_links_part_returns_empty(self):
        """Workbook without xl/externalLinks/ returns empty immediately."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Main"
        ws["A1"] = "='[1]Sheet1'!$A$1"
        # No external link attached
        result = self._save_and_check(wb, "no_ext_links.xlsx")
        self.assertEqual(result, [])


if __name__ == "__main__":
    unittest.main()
