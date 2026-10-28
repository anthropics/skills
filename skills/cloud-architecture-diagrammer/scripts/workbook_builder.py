"""Build the fixed workbook contract with editable inputs and live formulas."""

from io import BytesIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook
from openpyxl.drawing.image import Image as WorkbookImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from .model import CONTRACT, EstimateModel, ROW_KEYS, SCENARIOS
from .diagram_renderer import render_diagram

HEADER_ROW = 4
FIRST_DATA_ROW = HEADER_ROW + 1
BLUE, BLACK, GREEN = "0000FF", "000000", "008000"
NUMBER_FORMAT = '#,##0.00;(#,##0.00);"-"'


def _sheet(wb, name, headers, rows, table_name, explanation):
    ws = wb.create_sheet(name)
    ws.cell(1, 1, name.split("_", 1)[1]).font = Font(name="Arial", size=16, bold=True, color="17365D")
    ws.cell(2, 1, explanation).font = Font(name="Arial", color=BLACK)
    ws.cell(2, 1).alignment = Alignment(wrap_text=True)
    ws.row_dimensions[2].height = 32
    for column, header in enumerate(headers, 1):
        cell = ws.cell(HEADER_ROW, column, header)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(column)].width = min(45, max(16, len(header) + 3))
    for row_index, values in enumerate(rows or [[None] * len(headers)], FIRST_DATA_ROW):
        for column, value in enumerate(values, 1):
            cell = ws.cell(row_index, column, value)
            if isinstance(value, str):
                cell.data_type = "s"  # Brief text is data, even when it begins with '='.
            cell.font = Font(name="Arial", color=BLUE)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(value, (float, int)):
                cell.number_format = NUMBER_FORMAT
    last = max(FIRST_DATA_ROW, FIRST_DATA_ROW + len(rows) - 1)
    table = Table(displayName=table_name, ref=f"A{HEADER_ROW}:{get_column_letter(len(headers))}{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = f"A{FIRST_DATA_ROW}"
    ws.sheet_view.showGridLines = False
    ws.print_title_rows = f"1:{HEADER_ROW}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    return ws


def _formula(ws, row, column, formula, cross_sheet=False):
    cell = ws.cell(row, column, formula)
    cell.font = Font(name="Arial", color=GREEN if cross_sheet else BLACK)
    cell.number_format = NUMBER_FORMAT


def _validation(ws, column, last_row, allow_blank=False):
    validation = DataValidation(type="list", formula1='"' + ",".join(SCENARIOS) + '"', allow_blank=allow_blank)
    validation.errorTitle = "Scenario invalide"
    validation.error = "Choisir un scenario de la liste."
    validation.showErrorMessage = True
    validation.errorStyle = "stop"
    ws.add_data_validation(validation)
    validation.add(f"{column}{FIRST_DATA_ROW}:{column}{max(FIRST_DATA_ROW, last_row)}")


def _embed_architecture(ws, model, architecture, icon_pack):
    """Keep decision rows and the fixed header contract; append diagram records."""
    if architecture is None:
        architecture = {"nodes": [dict(row) for row in model.inventory], "edges": []}
    inventory_ids = [row["id"] for row in model.inventory]
    with TemporaryDirectory(prefix="estimate-diagram-") as directory:
        directory = Path(directory)
        manifest = render_diagram(architecture, directory / "architecture.svg", directory / "architecture.png", icon_pack)
        if len(inventory_ids) != len(set(inventory_ids)) or set(manifest.node_ids) != set(inventory_ids):
            raise ValueError("diagram node IDs must match workbook inventory IDs exactly")
        image = WorkbookImage(BytesIO((directory / "architecture.png").read_bytes()))
    row = ws.max_row + 3

    def records(title, headers, values):
        nonlocal row
        ws.cell(row, 1, title).font = Font(name="Arial", bold=True, color="17365D")
        row += 1
        for column, header in enumerate(headers, 1):
            ws.cell(row, column, header).font = Font(name="Arial", bold=True, color="17365D")
        row += 1
        for record in values:
            for column, value in enumerate(record, 1):
                cell = ws.cell(row, column, value)
                if isinstance(value, str):
                    cell.data_type = "s"
                cell.font = Font(name="Arial", color=BLACK)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            row += 1
        row += 2

    records("Diagram nodes", ["ID", "Composant", "Environnement", "Zone", "Fournisseur"],
            [[node["id"], node.get("component", node.get("label", node["id"])), node.get("environment", "Unspecified"), node.get("zone", "Unspecified"), node.get("provider", "")] for node in sorted(architecture["nodes"], key=lambda node: node["id"])])
    records("Diagram flows", ["ID", "Source", "Destination", "Flux"],
            [[edge["id"], edge["source"], edge["target"], edge.get("label", "")] for edge in sorted(architecture.get("edges", []), key=lambda edge: edge["id"])])
    records("Icon provenance", ["Fournisseur", "Source officielle", "Licence / marque", "Fichier", "SHA-256"],
            [[item["provider"], item["official_source_url"], item["licence_or_trademark_notice"], item["path"], item["sha256"]] for item in manifest.icon_provenance] or [["provider-neutral", "", "Original neutral labelled symbols", "", ""]])
    manifest_json = json.dumps(manifest.to_dict(), sort_keys=True, ensure_ascii=False)
    if len(manifest_json) > 32767:
        raise ValueError("diagram manifest exceeds Excel cell text limit")
    ws.cell(row, 1, "Diagram manifest").font = Font(name="Arial", bold=True, color="17365D")
    ws.cell(row + 1, 1, "JSON")
    ws.cell(row + 1, 2, manifest_json).data_type = "s"
    row += 4
    image.width, image.height = (image.width * min(1, 1100 / image.width), image.height * min(1, 1100 / image.width))
    ws.add_image(image, f"A{row}")


def build_workbook(model: EstimateModel, output: Path, *, architecture: dict | None = None, icon_pack: Path | None = None) -> Path:
    """Write twelve sections, formulas, and a self-contained inventory diagram.

    Optional architecture supplies explicit zones/flows; its node IDs must equal
    inventory IDs. Without it, draw inventory nodes and do not infer connections.
    """
    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    currency = model.metadata["currency"]
    sheets = CONTRACT["sheets"]
    _sheet(wb, sheets[0], ["Cle", "Valeur"], list(model.metadata.items()), "LisezMoi",
           "Donnees generiques. Effort en JH ; Build ponctuel et Run en JH/mois. Cout par fournisseur alternatif ; ne pas additionner les alternatives. Bleu : saisies, noir : calculs, vert : liens, jaune : hypotheses a confirmer.")
    assumptions = _sheet(wb, sheets[1], ["Hypothese", "Valeur", "Unite", "Statut", "Source", "Notes"],
                         [[row[key] for key in ROW_KEYS["assumptions"]] for row in model.assumptions], "Hypotheses",
                         "Scenario global et facteurs editables. Les prix synthetiques doivent etre remplaces par des sources datees.")
    assumption_rows = {}
    for row_index, item in enumerate(model.assumptions, FIRST_DATA_ROW):
        assumption_rows[item["key"]] = row_index
        if item["status"] == "A confirmer":
            for cell in assumptions[row_index]:
                cell.fill = PatternFill("solid", fgColor="FFFF00")
    for name in ("scenario_global", "taux_journalier", "mois_par_an"):
        row_index = assumption_rows[name]
        wb.defined_names.add(DefinedName(name, attr_text=f"'{assumptions.title}'!$B${row_index}"))
    scenario_validation = DataValidation(type="list", formula1='"' + ",".join(SCENARIOS) + '"')
    scenario_validation.showErrorMessage = True
    scenario_validation.showDropDown = False
    scenario_validation.errorStyle = "stop"
    scenario_validation.errorTitle = "Scenario invalide"
    scenario_validation.error = "Choisir Economique, Standard ou Resilient."
    assumptions.add_data_validation(scenario_validation)
    for key in ("scenario_global", "comparison_scenario"):
        if key in assumption_rows:
            scenario_validation.add(assumptions.cell(assumption_rows[key], 2))
    _sheet(wb, sheets[2], ["ID", "Composant", "Environnement", "Quantite", "Unite", "Notes"],
           [[row[key] for key in ROW_KEYS["inventory"]] for row in model.inventory], "Inventaire", "Inventaire generique des composants et environnements.")
    architecture_sheet = _sheet(wb, sheets[3], ["ID", "Composant", "Type", "Environnement", "Fournisseur", "Decision", "Notes"],
           [[row[key] for key in ROW_KEYS["architecture"]] for row in model.architecture], "Architecture", "Options d'architecture et decisions techniques a justifier.")
    _embed_architecture(architecture_sheet, model, architecture, icon_pack)
    task_headers = CONTRACT["task_columns"]
    task_columns = {header: index for index, header in enumerate(task_headers, 1)}
    task_rows = []
    for task in model.tasks:
        task_rows.append([task.id, task.lot, task.domain, task.task, task.deliverable, task.environment, task.role,
                          task.prerequisites, ", ".join(task.dependencies), task.acceptance_criterion, task.expected_evidence,
                          task.jh_low, task.jh_likely, task.jh_high, None, task.scenario_override, None, task.nature,
                          task.coverage, task.source, task.confidence, task.notes])
    tasks = _sheet(wb, sheets[4], task_headers, task_rows, "Taches", "Estimations en JH. Economique = bas, Standard = probable, Resilient = haut. Override local prioritaire ; Build et Run se lisent separement.")
    for row_index, task in enumerate(model.tasks, FIRST_DATA_ROW):
        ref = {name: f"{get_column_letter(column)}{row_index}" for name, column in task_columns.items()}
        _formula(tasks, row_index, task_columns["Scenario retenu"], f'=IF({ref["Scenario override"]}="",scenario_global,{ref["Scenario override"]})', True)
        _formula(tasks, row_index, task_columns["JH retenu"], f'=IF({ref["Scenario retenu"]}="{SCENARIOS[0]}",{ref["JH bas"]},IF({ref["Scenario retenu"]}="{SCENARIOS[1]}",{ref["JH probable"]},{ref["JH haut"]}))')
    _validation(tasks, get_column_letter(task_columns["Scenario override"]), FIRST_DATA_ROW + len(model.tasks) - 1, True)
    cost_headers = ["ID", "Fournisseur", "Service", "Environnement", "Scenario", "Unite", "Quantite", f"Prix unitaire ({currency})", "Devise", "Multiplicateur mensuel", f"Cout mensuel ({currency})", f"Cout annuel ({currency})", "Source", "Confiance", "Notes", "Statut prix"]
    costs = _sheet(wb, sheets[5], cost_headers,
                   [[cost.id, cost.provider, cost.service, cost.environment, cost.scenario, cost.unit, cost.quantity,
                     cost.unit_price, cost.currency, cost.monthly_multiplier, None, None, cost.source, cost.confidence,
                     cost.notes, "A confirmer"] for cost in model.cloud_costs], "CoutsCloud", "Coûts recurrents par option fournisseur et scenario ; tarifs synthetiques non contractuels. Le multiplicateur convertit l'unite en mois.")
    cost_columns = {header: get_column_letter(index) for index, header in enumerate(cost_headers, 1)}
    for row_index, cost in enumerate(model.cloud_costs, FIRST_DATA_ROW):
        quantity = cost_columns["Quantite"]
        price = cost_columns[f"Prix unitaire ({currency})"]
        factor = cost_columns["Multiplicateur mensuel"]
        monthly = cost_columns[f"Cout mensuel ({currency})"]
        _formula(costs, row_index, cost_headers.index(f"Cout mensuel ({currency})") + 1, f"={quantity}{row_index}*{price}{row_index}*{factor}{row_index}")
        _formula(costs, row_index, cost_headers.index(f"Cout annuel ({currency})") + 1, f"={monthly}{row_index}*mois_par_an", cross_sheet=True)
    scenario_rows = [[scenario.name, scenario.description, None, None, None, scenario.notes] for scenario in model.scenarios]
    scenarios = _sheet(wb, sheets[6], ["Scenario", "Description", "Rapprochement technique (non budgetaire)", "JH Build", "JH Run/mois", "Notes"], scenario_rows, "Scenarios", "Comparaison sans overrides locaux. Le rapprochement technique additionne JH Build et JH Run/mois, des unites incompatibles : utilisation budgetaire interdite.")
    effort_columns = dict(zip(SCENARIOS, ("JH bas", "JH probable", "JH haut")))
    for row_index, scenario in enumerate(model.scenarios, FIRST_DATA_ROW):
        column = effort_columns[scenario.name]
        _formula(scenarios, row_index, 3, f"=SUM(Taches[{column}])", True)
        _formula(scenarios, row_index, 4, f'=SUMIF(Taches[Nature],"Build",Taches[{column}])', True)
        _formula(scenarios, row_index, 5, f'=SUMIF(Taches[Nature],"Run",Taches[{column}])', True)
    phases = _sheet(wb, sheets[7], ["Phase", "Sequence", "Jalon", "Rapprochement technique (non budgetaire)", "JH Build", "JH Run/mois", "Notes"],
                    [[row["phase"], row["sequence"], row["milestone"], None, None, None, row["notes"]] for row in model.phasing], "Phasage", "Lots et jalons. Le rapprochement technique additionne JH Build et JH Run/mois, des unites incompatibles : utilisation budgetaire interdite.")
    for row_index, phase in enumerate(model.phasing, FIRST_DATA_ROW):
        _formula(phases, row_index, 4, f"=SUMIF(Taches[Lot],A{row_index},Taches[JH retenu])", True)
        for column, nature in ((5, "Build"), (6, "Run")):
            _formula(phases, row_index, column, f'=SUMIFS(Taches[JH retenu],Taches[Lot],A{row_index},Taches[Nature],"{nature}")', True)
    _sheet(wb, sheets[8], ["ID", "Risque", "Impact", "Probabilite", "Mitigation", "Responsable", "Statut"],
           [[row[key] for key in ROW_KEYS["risks"]] for row in model.risks], "Risques", "Incertitudes, impacts et actions de reduction du risque.")
    _sheet(wb, sheets[9], ["ID", "Question", "Option", "Justification", "Statut", "Source"],
           [[row[key] for key in ROW_KEYS["decisions"]] for row in model.decisions], "Arbitrages", "Decisions explicites et hypotheses a confirmer.")
    _sheet(wb, sheets[10], ["ID", "Titre", "Type", "Date", "URL", "Notes"],
           [[row[key] for key in ROW_KEYS["sources"]] for row in model.sources], "Sources", "References datees ; les exemples synthetiques ne constituent pas des tarifs de marche.")
    summary_rows = [["Rapprochement technique (Build + Run/mois)", None, "Unites mixtes", "Utilisation budgetaire interdite : JH Build et JH Run/mois ont des unites incompatibles."],
                    ["JH Build", None, "JH", "Effort ponctuel."], ["JH Run mensuel", None, "JH/mois", "Effort recurrent."],
                    ["Cout Build", None, currency, "Effort Build au taux journalier saisi."],
                    ["Cout Run mensuel", None, f"{currency}/mois", "Effort Run au taux journalier saisi."]]
    providers = sorted({cost.provider for cost in model.cloud_costs})
    for provider in providers:
        for scenario in SCENARIOS:
            summary_rows.append([f"Cloud mensuel | {provider} | {scenario}", None, f"{currency}/mois", "Option alternative ; ne pas sommer les fournisseurs."])
    summary_rows.append(["Statut finance", None, "Controle", "Un taux inconnu bloque couts humains, TCO et approbation financiere, pas les capacites ou JH. Un taux renseigne ne vaut pas approbation."])
    summary = _sheet(wb, sheets[11], ["Indicateur", "Valeur", "Unite", "Notes"], summary_rows, "Synthese", "Synthese dynamique des charges et options cloud. Les prix, la couverture et les incertitudes doivent etre confirmes avant engagement.")
    summary_formulas = ["=SUM(Taches[JH retenu])", '=SUMIF(Taches[Nature],"Build",Taches[JH retenu])',
                        '=SUMIF(Taches[Nature],"Run",Taches[JH retenu])',
                        '=IF(taux_journalier="","",IF(taux_journalier="A chiffrer","",SUMIF(Taches[Nature],"Build",Taches[JH retenu])*taux_journalier))',
                        '=IF(taux_journalier="","",IF(taux_journalier="A chiffrer","",SUMIF(Taches[Nature],"Run",Taches[JH retenu])*taux_journalier))']
    for provider in providers:
        for scenario in SCENARIOS:
            quoted_provider = provider.replace('"', '""')
            summary_formulas.append(f'=SUMIFS(CoutsCloud[Cout mensuel ({currency})],CoutsCloud[Fournisseur],"{quoted_provider}",CoutsCloud[Scenario],"{scenario}")')
    summary_formulas.append('=IF(taux_journalier="","Bloque - taux journalier A chiffrer",IF(taux_journalier="A chiffrer","Bloque - taux journalier A chiffrer","Taux journalier renseigne - approbation non acquise"))')
    for row_index, formula in enumerate(summary_formulas, FIRST_DATA_ROW):
        _formula(summary, row_index, 2, formula, True)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return output
