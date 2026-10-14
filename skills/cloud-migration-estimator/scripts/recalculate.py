"""Recalculate atomically with a native engine or a calibrated, limited evaluator.

The Python engine supports this suite's arithmetic, references, IF, SUM, SUMIF,
and SUMIFS. Unknown syntax fails closed; it is not a general Excel replacement.
"""

from dataclasses import asdict, dataclass, field
import math
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import range_boundaries
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table

from .publication import blank_formula_caches, sanitize_publication

ERROR_TOKENS = ("#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


@dataclass
class RecalculationResult:
    ok: bool = False
    engine: str = "none"
    formula_count: int = 0
    cached_values_count: int = 0
    controls_passed: int = 0
    native_opened: bool = False
    publication_inspected_parts: int = 0
    final_native_opened: bool = False
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def to_dict(self):
        return asdict(self)


class FormulaFailure(ValueError):
    pass


class _Parser:
    PRECEDENCE = {"=": 1, "<>": 1, "<": 1, ">": 1, "<=": 1, ">=": 1,
                  "+": 2, "-": 2, "*": 3, "/": 3, "^": 4}

    def __init__(self, formula):
        self.tokens = [token for token in Tokenizer(formula).items if token.type != "WHITE-SPACE"]
        self.position = 0

    def take(self):
        if self.position >= len(self.tokens):
            raise FormulaFailure("unsupported incomplete formula")
        token = self.tokens[self.position]
        self.position += 1
        return token

    def expression(self, minimum=0):
        token = self.take()
        if token.type == "OPERATOR-PREFIX" and token.value in ("+", "-"):
            node = ("unary", token.value, self.expression(5))
        elif token.type == "FUNC" and token.subtype == "OPEN":
            arguments = []
            if self.tokens[self.position].subtype != "CLOSE":
                while True:
                    arguments.append(self.expression())
                    if self.tokens[self.position].type != "SEP":
                        break
                    self.take()
            if self.take().subtype != "CLOSE":
                raise FormulaFailure("unsupported function syntax")
            node = ("function", token.value[:-1].upper(), arguments)
        elif token.type == "PAREN" and token.subtype == "OPEN":
            node = self.expression()
            if self.take().subtype != "CLOSE":
                raise FormulaFailure("unsupported parentheses")
        elif token.type == "OPERAND":
            node = ("operand", token.subtype, token.value)
        else:
            raise FormulaFailure(f"unsupported token {token.value}")
        while self.position < len(self.tokens):
            token = self.tokens[self.position]
            precedence = self.PRECEDENCE.get(token.value, -1)
            if token.type != "OPERATOR-INFIX" or precedence < minimum:
                break
            self.take()
            node = ("binary", token.value, node, self.expression(precedence + 1))
        return node

    def parse(self):
        node = self.expression()
        if self.position != len(self.tokens):
            raise FormulaFailure("unsupported trailing syntax")
        return node


def _number(value):
    if value is None or value == "":
        return 0
    if not isinstance(value, (int, float, bool)):
        raise FormulaFailure("#VALUE! nonnumeric arithmetic")
    return value


def _flatten(value):
    return value if isinstance(value, list) else [value]


def _criteria_text(value):
    return "" if value is None else str(value).casefold()


class _Evaluator:
    def __init__(self, workbook):
        self.workbook = workbook
        self.cache = {}
        self.active = set()

    def cell(self, sheet, coordinate):
        key = (sheet, coordinate.replace("$", ""))
        if key in self.active:
            raise FormulaFailure("circular reference")
        if key in self.cache:
            return self.cache[key]
        if sheet not in self.workbook.sheetnames:
            raise FormulaFailure(f"#REF! missing sheet {sheet}")
        cell = self.workbook[sheet][key[1]]
        self.active.add(key)
        try:
            value = self.node(_Parser(cell.value).parse(), sheet) if cell.data_type == "f" else cell.value
            if cell.data_type == "e":
                raise FormulaFailure(str(value))
            if isinstance(value, float) and not math.isfinite(value):
                raise FormulaFailure("#NUM! nonfinite result")
            self.cache[key] = value
            return value
        finally:
            self.active.remove(key)

    def reference(self, reference, sheet):
        structured = re.fullmatch(r"(\w+)\[(.+)\]", reference)
        if structured:
            table_name, header = structured.groups()
            for ws in self.workbook:
                if table_name in ws.tables:
                    table = ws.tables[table_name]
                    left, top, right, bottom = range_boundaries(table.ref)
                    headers = [ws.cell(top, column).value for column in range(left, right + 1)]
                    if header not in headers:
                        raise FormulaFailure(f"#REF! missing table column {reference}")
                    column = left + headers.index(header)
                    return [self.cell(ws.title, ws.cell(row, column).coordinate) for row in range(top + 1, bottom + 1)]
            raise FormulaFailure(f"#REF! missing table {table_name}")
        if reference in self.workbook.defined_names:
            destinations = list(self.workbook.defined_names[reference].destinations)
            if len(destinations) != 1:
                raise FormulaFailure("unsupported named reference")
            target_sheet, coordinate = destinations[0]
            return self.reference(coordinate, target_sheet)
        if "!" in reference:
            sheet, reference = reference.rsplit("!", 1)
            sheet = sheet.strip("'").replace("''", "'")
        if not re.fullmatch(r"\$?[A-Za-z]{1,3}\$?\d+(?::\$?[A-Za-z]{1,3}\$?\d+)?", reference):
            raise FormulaFailure(f"unsupported reference {reference}")
        if ":" in reference:
            left, top, right, bottom = range_boundaries(reference)
            return [self.cell(sheet, f"{self.workbook[sheet].cell(row, col).coordinate}")
                    for row in range(top, bottom + 1) for col in range(left, right + 1)]
        return self.cell(sheet, reference)

    def node(self, node, sheet):
        kind = node[0]
        if kind == "operand":
            _, subtype, value = node
            if subtype == "NUMBER":
                return float(value)
            if subtype == "TEXT":
                return value[1:-1].replace('""', '"')
            if subtype == "LOGICAL":
                return value == "TRUE"
            if subtype == "ERROR":
                raise FormulaFailure(value)
            return self.reference(value, sheet)
        if kind == "unary":
            value = _number(self.node(node[2], sheet))
            return value if node[1] == "+" else -value
        if kind == "binary":
            _, operator, first, second = node
            left, right = self.node(first, sheet), self.node(second, sheet)
            if operator in ("=", "<>", "<", ">", "<=", ">="):
                left, right = ("" if left is None else left), ("" if right is None else right)
                if operator == "=":
                    return left == right
                if operator == "<>":
                    return left != right
                return {"<": lambda: left < right, ">": lambda: left > right,
                        "<=": lambda: left <= right, ">=": lambda: left >= right}[operator]()
            left, right = _number(left), _number(right)
            if operator == "/" and right == 0:
                raise FormulaFailure("#DIV/0!")
            return {"+": lambda: left + right, "-": lambda: left - right,
                    "*": lambda: left * right, "/": lambda: left / right,
                    "^": lambda: left ** right}[operator]()
        _, function, arguments = node
        if function == "IF" and len(arguments) == 3:
            return self.node(arguments[1 if self.node(arguments[0], sheet) else 2], sheet)
        if function not in ("SUM", "SUMIF", "SUMIFS"):
            raise FormulaFailure(f"unsupported function {function}")
        values = [self.node(argument, sheet) for argument in arguments]
        if function == "SUM":
            return sum(value for item in values for value in _flatten(item) if isinstance(value, (int, float)))
        if function == "SUMIF" and len(values) in (2, 3):
            criteria = [(_flatten(values[0]), values[1])]
            numbers = _flatten(values[2] if len(values) == 3 else values[0])
        elif function == "SUMIFS" and len(values) >= 3 and len(values) % 2 == 1:
            numbers = _flatten(values[0])
            criteria = [(_flatten(values[index]), values[index + 1]) for index in range(1, len(values), 2)]
        else:
            raise FormulaFailure(f"unsupported arguments for {function}")
        if any(len(items) != len(numbers) for items, _ in criteria):
            raise FormulaFailure("#VALUE! range lengths differ")
        # This suite uses literal equality criteria only. Fail closed for other Excel criteria.
        if any(isinstance(expected, str) and re.search(r"[?*<>=~]", expected) for _, expected in criteria):
            raise FormulaFailure("unsupported comparison or wildcard criterion")
        return sum(_number(value) for index, value in enumerate(numbers)
                   if all(_criteria_text(items[index]) == _criteria_text(expected) for items, expected in criteria))


def _calibrate_python():
    """Control expectations are hand-derived and differ from realistic defects."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Control"
    for row in [("Amount", "Group", "Environment"), (2, "A", "Prod"), (7, "A", "Dev"), (11, "B", "Prod")]:
        ws.append(row)
    ws.add_table(Table(displayName="Controls", ref="A1:C4"))
    wb.defined_names.add(DefinedName("factor", attr_text="'Control'!$A$2"))
    controls = [
        ("E1", "=A2*3", 6, "F1", "=A2*0", 0),
        ("E2", "=SUM(E1,A3)", 13, "F2", "=SUM(A3)", 7),
        ("E3", '=SUMIFS(Controls[Amount],Controls[Group],"A",Controls[Environment],"Prod")',
         2, "F3", '=SUMIF(Controls[Group],"A",Controls[Amount])', 9),
        ("E4", '=IF(A2=2,factor*5,1/0)', 10, "F4", '=IF(A2=0,factor*5,0)', 0),
        ("E5", '=SUMIF(Controls[Group],"B",Controls[Amount])', 11,
         "F5", '=SUM(Controls[Amount])', 20),
    ]
    for correct_cell, correct_formula, _, broken_cell, broken_formula, _ in controls:
        ws[correct_cell] = correct_formula
        ws[broken_cell] = broken_formula
    evaluator = _Evaluator(wb)
    errors = []
    passed = 0
    for coordinate, _, expected, broken_coordinate, _, defective in controls:
        try:
            actual = evaluator.cell("Control", coordinate)
            if actual == expected:
                passed += 1
            else:
                errors.append(f"calibration control {coordinate} returned {actual!r}, expected {expected!r}")
            broken_actual = evaluator.cell("Control", broken_coordinate)
            if expected != defective and broken_actual == defective:
                passed += 1
            else:
                errors.append(f"calibration defect {broken_coordinate} returned {broken_actual!r}, expected {defective!r} distinct from {expected!r}")
        except Exception as error:
            errors.append(f"calibration control {coordinate}: {error}")
    return passed, errors


def _run_excel_powershell(script, path=None, timeout=120):
    executable = shutil.which("powershell.exe") or shutil.which("pwsh.exe")
    if executable is None:
        raise FormulaFailure("PowerShell is unavailable for Excel COM")
    environment = os.environ.copy()
    if path is not None:
        environment["ESTIMATE_RECALC_PATH"] = str(path.resolve())
    # Pass an argument vector directly (no shell expansion). EncodedCommand can
    # terminate during COM Open with exit 0; require a positive completion signal.
    marker = "ESTIMATE_RECALC_COMPLETED"
    command = _EXCEL_HELPERS + script + f'\nWrite-Output "{marker}"\n'
    result = subprocess.run([executable, "-NoProfile", "-NonInteractive", "-Command", command],
                            capture_output=True, text=True, timeout=timeout, env=environment)
    if result.returncode or marker not in result.stdout.splitlines():
        raise FormulaFailure(f"PowerShell Excel COM failed: {(result.stderr or result.stdout).strip()}")


_EXCEL_HELPERS = """
function Release-ExcelComObject {
    param($object)
    [void][System.Runtime.InteropServices.Marshal]::FinalReleaseComObject($object)
}
function Invoke-ExcelCollection { [GC]::Collect() }
function Wait-ExcelFinalizers { [GC]::WaitForPendingFinalizers() }
function Close-ExcelResources {
    param($workbook, $excel_workbooks, $excel)
    $cleanup_failures = [Collections.Generic.List[string]]::new()
    if ($null -ne $workbook) {
        try { $workbook.Close($false) }
        catch { $cleanup_failures.Add('Workbook.Close: ' + $_.Exception.Message) }
        try { Release-ExcelComObject $workbook }
        catch { $cleanup_failures.Add('Workbook release: ' + $_.Exception.Message) }
    }
    if ($null -ne $excel_workbooks) {
        try { Release-ExcelComObject $excel_workbooks }
        catch { $cleanup_failures.Add('Workbooks release: ' + $_.Exception.Message) }
    }
    if ($null -ne $excel) {
        try { $excel.Quit() }
        catch { $cleanup_failures.Add('Excel.Quit: ' + $_.Exception.Message) }
        try { Release-ExcelComObject $excel }
        catch { $cleanup_failures.Add('Excel release: ' + $_.Exception.Message) }
    }
    try { Invoke-ExcelCollection }
    catch { $cleanup_failures.Add('GC collection: ' + $_.Exception.Message) }
    try { Wait-ExcelFinalizers }
    catch { $cleanup_failures.Add('GC finalizers: ' + $_.Exception.Message) }
    return $cleanup_failures.ToArray()
}
"""


_EXCEL_PROBE = """
$ErrorActionPreference = 'Stop'
$excel = $null
$failures = [Collections.Generic.List[string]]::new()
try {
    $excel = New-Object -ComObject Excel.Application
} catch {
    $failures.Add('Excel probe: ' + $_.Exception.Message)
} finally {
    foreach ($failure in @(Close-ExcelResources $null $null $excel)) {
        $failures.Add($failure)
    }
}
if ($failures.Count -gt 0) { throw ($failures -join '; ') }
"""


_EXCEL_RECALCULATE = """
$ErrorActionPreference = 'Stop'
$excel = $null
$workbook = $null
$excel_workbooks = $null
$failures = [Collections.Generic.List[string]]::new()
try {
    $excel = New-Object -ComObject Excel.Application
    $excel.Visible = $false
    $excel.DisplayAlerts = $false
    $excel.AutomationSecurity = 3
    $excel_workbooks = $excel.Workbooks
    $workbook = $excel_workbooks.Open($env:ESTIMATE_RECALC_PATH, 0, $false)
    $excel.CalculateFullRebuild()
    $workbook.Save()
} catch {
    $failures.Add('Excel recalculation: ' + $_.Exception.Message)
} finally {
    foreach ($failure in @(Close-ExcelResources $workbook $excel_workbooks $excel)) {
        $failures.Add($failure)
    }
}
if ($failures.Count -gt 0) { throw ($failures -join '; ') }
"""


def _select_excel_backend():
    if sys.platform != "win32":
        return None
    try:
        import win32com.client
        excel = win32com.client.DispatchEx("Excel.Application")
        try:
            return "pywin32"
        finally:
            excel.Quit()
    except Exception:
        pass
    try:
        _run_excel_powershell(_EXCEL_PROBE, timeout=30)
        return "powershell"
    except Exception:
        return None


def _find_libreoffice():
    for name in ("soffice", "soffice.exe", "libreoffice"):
        found = shutil.which(name)
        if found:
            return Path(found)
    if sys.platform == "win32":
        for root in (os.environ.get("ProgramFiles", r"C:\Program Files"), os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")):
            candidate = Path(root) / "LibreOffice/program/soffice.exe"
            if candidate.is_file():
                return candidate
    return None


def _close_excel_pywin32(workbook, excel):
    """Each owned COM resource receives its attempt even if cleanup fails."""
    failures = []
    if workbook is not None:
        try:
            workbook.Close(SaveChanges=False)
        except Exception as error:
            failures.append(f"Workbook.Close: {error}")
    try:
        excel.Quit()
    except Exception as error:
        failures.append(f"Excel.Quit: {error}")
    return failures


def _native_excel(path, backend):
    if backend == "powershell":
        _run_excel_powershell(_EXCEL_RECALCULATE, path)
        return
    if backend != "pywin32":
        raise FormulaFailure(f"Unknown Excel COM backend: {backend}")
    import win32com.client
    excel = win32com.client.DispatchEx("Excel.Application")
    workbook = None
    failures = []
    try:
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.AutomationSecurity = 3  # Disable macros for the workbook being inspected.
        workbook = excel.Workbooks.Open(str(path.resolve()), UpdateLinks=0, ReadOnly=False)
        excel.CalculateFullRebuild()
        workbook.Save()
    except Exception as error:
        failures.append(f"Excel recalculation: {error}")
    finally:
        failures.extend(_close_excel_pywin32(workbook, excel))
    if failures:
        raise FormulaFailure("; ".join(failures))


def _native_libreoffice(path, executable, directory):
    output = directory / "converted"
    output.mkdir()
    profile = (directory / "profile").as_uri()
    result = subprocess.run([str(executable), f"-env:UserInstallation={profile}", "--headless", "--convert-to",
                             "xlsx:Calc MS Excel 2007 XML", "--outdir", str(output), str(path)],
                            capture_output=True, text=True, timeout=120)
    converted = output / path.name
    if result.returncode or not converted.is_file():
        raise FormulaFailure(f"LibreOffice conversion failed: {result.stderr.strip()}")
    shutil.copyfile(converted, path)


def _verify_native_final(path, backend):
    """Open sanitized bytes read-only, close without Save, and assert unchanged."""
    before = Path(path).read_bytes()
    if backend == "powershell":
        script = _EXCEL_RECALCULATE.replace("$false)", "$true)", 1).replace("    $excel.CalculateFullRebuild()\n    $workbook.Save()", "    if (-not $workbook.ReadOnly) { throw 'Final workbook is not read-only' }")
        _run_excel_powershell(script, path)
    else:
        import win32com.client
        excel = win32com.client.DispatchEx("Excel.Application")
        workbook = None
        failures = []
        try:
            excel.Visible = False
            excel.DisplayAlerts = False
            excel.AutomationSecurity = 3
            workbook = excel.Workbooks.Open(str(Path(path).resolve()), UpdateLinks=0, ReadOnly=True)
            if not workbook.ReadOnly:
                raise FormulaFailure("Final workbook is not read-only")
        except Exception as error:
            failures.append(f"Final Excel opening: {error}")
        finally:
            failures.extend(_close_excel_pywin32(workbook, excel))
        if failures:
            raise FormulaFailure("; ".join(failures))
    if Path(path).read_bytes() != before:
        raise FormulaFailure("Final native opening changed published bytes")


def _formula_cells(workbook):
    return [(ws.title, cell.coordinate) for ws in workbook for row in ws for cell in row if cell.data_type == "f"]


def _inject_caches(source, destination, workbook, values):
    # Edit only worksheet XML; styles, tables, formulas, and workbook metadata survive.
    with ZipFile(source) as archive:
        relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {item.attrib["Id"]: item.attrib["Target"] for item in relationships}
        metadata = ET.fromstring(archive.read("xl/workbook.xml"))
        sheet_paths = {}
        for sheet in metadata.find(f"{{{NS}}}sheets"):
            target = targets[sheet.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]]
            sheet_paths[target.lstrip("/") if target.startswith("/") else "xl/" + target] = sheet.attrib["name"]
        with ZipFile(destination, "w") as result:
            for item in archive.infolist():
                data = archive.read(item.filename)
                if item.filename in sheet_paths:
                    root = ET.fromstring(data)
                    name = sheet_paths[item.filename]
                    for cell in root.iter(f"{{{NS}}}c"):
                        if cell.find(f"{{{NS}}}f") is None:
                            continue
                        value = values[(name, cell.attrib["r"])]
                        cached = cell.find(f"{{{NS}}}v")
                        if cached is None:
                            cached = ET.SubElement(cell, f"{{{NS}}}v")
                        if isinstance(value, str) or value is None:
                            cell.set("t", "str")
                            cached.text = "" if value is None else value
                        elif isinstance(value, bool):
                            cell.set("t", "b")
                            cached.text = str(int(value))
                        else:
                            cell.pop("t", None) if hasattr(cell, "pop") else cell.attrib.pop("t", None)
                            cached.text = str(value)
                    data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                result.writestr(item, data)


def recalculate(path: Path) -> RecalculationResult:
    """Commit recalculated bytes only after a values-only reread succeeds."""
    path = Path(path).resolve()
    report = RecalculationResult()
    try:
        # Embedded images can retain a ZipExtFile beyond openpyxl's archive close.
        # Own the stream lifetime so Windows allows the atomic replacement.
        with path.open("rb") as stream:
            workbook = load_workbook(stream)
        formulas = _formula_cells(workbook)
        report.formula_count = len(formulas)
        with tempfile.TemporaryDirectory(prefix="estimate-recalc-", dir=path.parent) as temporary:
            directory = Path(temporary)
            candidate = directory / path.name
            shutil.copyfile(path, candidate)
            engines = []
            excel_backend = _select_excel_backend()
            if excel_backend:
                engines.append(("excel-com", lambda: _native_excel(candidate, excel_backend)))
            libreoffice = _find_libreoffice()
            if libreoffice:
                engines.append(("libreoffice", lambda: _native_libreoffice(candidate, libreoffice, directory)))
            for engine, action in engines:
                try:
                    action()
                    _verify_caches(candidate, formulas)
                    report.engine, report.native_opened = engine, True
                    break
                except Exception as error:
                    report.warnings.append(f"{engine} unavailable or failed: {error}")
                    shutil.copyfile(path, candidate)
            else:
                report.engine = "python-calibrated"
                report.controls_passed, report.errors = _calibrate_python()
                if report.errors:
                    return report
                evaluator = _Evaluator(workbook)
                values = {}
                for sheet, coordinate in formulas:
                    try:
                        values[(sheet, coordinate)] = evaluator.cell(sheet, coordinate)
                    except Exception as error:
                        report.errors.append(f"{sheet}!{coordinate}: {error}")
                if report.errors:
                    return report
                _inject_caches(path, candidate, workbook, values)
                report.warnings.append("Python fallback used; native Excel/LibreOffice opening is not verified. Only the documented formula subset is supported.")
            report.publication_inspected_parts = sanitize_publication(candidate)
            _verify_caches(candidate, formulas)
            if report.engine == "excel-com":
                _verify_native_final(candidate, excel_backend)
                report.final_native_opened = True
            report.cached_values_count = len(formulas)
            os.replace(candidate, path)
            report.ok = True
    except Exception as error:
        report.errors.append(str(error))
    return report


def _verify_caches(path, formulas):
    blanks = blank_formula_caches(path)
    with Path(path).open("rb") as stream:
        values = load_workbook(stream, data_only=True)
    with Path(path).open("rb") as stream:
        formula_workbook = load_workbook(stream)
    for sheet, coordinate in formulas:
        if formula_workbook[sheet][coordinate].data_type != "f":
            raise FormulaFailure(f"{sheet}!{coordinate}: native conversion lost formula")
        cell = values[sheet][coordinate]
        if cell.value is None and (sheet, coordinate) not in blanks:
            raise FormulaFailure(f"{sheet}!{coordinate}: missing cached value")
        if cell.data_type == "e" or isinstance(cell.value, str) and cell.value in ERROR_TOKENS:
            raise FormulaFailure(f"{sheet}!{coordinate}: {cell.value}")
