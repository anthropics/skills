"""Read-only structural, publication, and independent financial validation."""

from collections import defaultdict
from dataclasses import asdict, dataclass, field
import json
import math
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from .recalculate import ERROR_TOKENS, _Evaluator, _calibrate_python
from .publication import blank_formula_caches

DEFAULT_CONTRACT = Path(__file__).with_name("workbook-contract.json")
FUTURE_FUNCTIONS = re.compile(r"(?<![\w.])(?:IFNA|IFS|XLOOKUP|XMATCH|CONCAT|TEXTJOIN|FILTER|UNIQUE|SORT|LET|LAMBDA)\s*\(", re.I)


@dataclass
class ValidationReport:
    ok: bool = False
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    finance_status: str = "not_checked"
    counts: dict[str, int] = field(default_factory=lambda: dict.fromkeys(
        ("sheets", "formulas", "formula_errors", "formula_semantic_errors", "missing_cached_values", "table_filter_overlaps",
         "unsupported_future_functions", "reconciliations", "reconciliation_errors", "diagram_inventory_ids", "private_data_matches", "annex_blocks", "errors"), 0))
    cost_totals: list[dict] = field(default_factory=list)

    def to_dict(self):
        return asdict(self)


def _overlap(first, second):
    a, b, c, d = range_boundaries(first)
    e, f, g, h = range_boundaries(second)
    return max(a, e) <= min(c, g) and max(b, f) <= min(d, h)


def _rows(workbook, name):
    ws = workbook[name]
    # Scenarios may also contain a separately validated provider comparison.
    # Other sheets still inspect every populated row, including outside tables.
    primary_tables = ws.tables.values() if name == "06_Scenarios" else ()
    bounds = next((range_boundaries(table.ref) for table in primary_tables
                   if range_boundaries(table.ref)[:2] == (1, 4)), (1, 4, ws.max_column, ws.max_row))
    left, top, right, bottom = bounds
    headers = [ws.cell(top, column).value for column in range(left, right + 1)]
    return [(index, dict(zip(headers, [cell.value for cell in ws[index]])))
            for index in range(top + 1, bottom + 1) if ws.cell(index, left).value is not None]


def _assumptions(workbook):
    return {row["Hypothese"]: row["Valeur"] for _, row in _rows(workbook, "01_Hypotheses")}


def _same_formula_value(actual, expected):
    if isinstance(expected, bool):
        return isinstance(actual, bool) and actual == expected
    if isinstance(expected, (int, float)):
        return (isinstance(actual, (int, float)) and not isinstance(actual, bool)
                and math.isfinite(actual) and math.isclose(actual, expected, rel_tol=1e-9, abs_tol=1e-7))
    return actual == expected


def validate_workbook(path: Path, contract: Path = DEFAULT_CONTRACT, *, asserted_totals=None,
                      diagram_inventory_ids=None) -> ValidationReport:
    """Do not recalculate or repair: stale/missing caches must fail this gate.

    Supply every numerical prose assertion as {"sheet!cell": expected_value}.
    Diagram IDs must match the inventory exactly when supplied.
    """
    report = ValidationReport()
    try:
        specification = json.loads(Path(contract).read_text(encoding="utf-8"))
        formulas = load_workbook(path)
        values = load_workbook(path, data_only=True)
        blanks = blank_formula_caches(path)
    except Exception as error:
        report.errors.append(f"Cannot inspect workbook/contract: {error}")
        report.counts["errors"] = len(report.errors)
        return report
    controls_passed, control_errors = _calibrate_python()
    if control_errors or controls_passed != 10:
        report.errors.append(f"Formula/cache evaluator calibration failed: {control_errors or controls_passed}")
    evaluator = _Evaluator(formulas)
    report.counts["sheets"] = len(formulas.sheetnames)
    for ws in formulas:
        for row in ws:
            for cell in row:
                location = f"{ws.title}!{cell.coordinate}"
                text = str(cell.value or "")
                cached = values[ws.title][cell.coordinate]
                tokens = [token for token in ERROR_TOKENS if token in text or cached.value == token]
                if tokens:
                    report.counts["formula_errors"] += 1
                    report.errors.append(f"{location}: formula error {', '.join(tokens)}")
                if cell.data_type == "f":
                    report.counts["formulas"] += 1
                    if cached.value is None and (ws.title, cell.coordinate) not in blanks:
                        report.counts["missing_cached_values"] += 1
                        report.errors.append(f"{location}: missing cached value; recalculate before delivery")
                    elif not control_errors and controls_passed == 10:
                        try:
                            evaluated = evaluator.cell(ws.title, cell.coordinate)
                            actual = "" if (ws.title, cell.coordinate) in blanks else cached.value
                            if not _same_formula_value(actual, evaluated):
                                raise ValueError(f"cached {cached.value!r}, evaluated {evaluated!r}")
                        except Exception as error:
                            report.counts["formula_semantic_errors"] += 1
                            report.errors.append(f"{location}: formula/cache mismatch ({error})")
                    if FUTURE_FUNCTIONS.search(text) or "_xlfn." in text:
                        report.counts["unsupported_future_functions"] += 1
                        report.errors.append(f"{location}: unsupported future function")
                elif isinstance(cell.value, str):
                    for pattern in specification.get("forbidden_patterns", []):
                        if re.search(pattern, cell.value):
                            report.counts["private_data_matches"] += 1
                            report.errors.append(f"{location}: private data pattern matched")
                            break
        tables = list(ws.tables.values())
        for index, table in enumerate(tables):
            if ws.auto_filter.ref and _overlap(table.ref, ws.auto_filter.ref):
                report.counts["table_filter_overlaps"] += 1
                report.errors.append(f"{ws.title}: table/filter overlap {table.ref} / {ws.auto_filter.ref}")
            for other in tables[index + 1:]:
                if _overlap(table.ref, other.ref):
                    report.errors.append(f"{ws.title}: overlapping tables {table.name} / {other.name}")
    if formulas.sheetnames != specification["sheets"]:
        report.errors.append(f"Sheet contract mismatch: expected {specification['sheets']}; found {formulas.sheetnames}")
    task_name = "04_Taches_JH"
    if task_name in formulas.sheetnames:
        headers = [cell.value for cell in formulas[task_name][4]]
        if headers != specification["task_columns"]:
            report.errors.append(f"{task_name}: column contract mismatch")
    expected_headers = {
        "00_Lisez-moi": ["Cle", "Valeur"],
        "01_Hypotheses": ["Hypothese", "Valeur", "Unite", "Statut", "Source", "Notes"],
        "02_Inventaire": ["ID", "Composant", "Environnement", "Quantite", "Unite", "Notes"],
        "03_Architecture": ["ID", "Composant", "Type", "Environnement", "Fournisseur", "Decision", "Notes"],
        "06_Scenarios": ["Scenario", "Description", "Rapprochement technique (non budgetaire)", "JH Build", "JH Run/mois", "Notes"],
        "07_Phasage": ["Phase", "Sequence", "Jalon", "Rapprochement technique (non budgetaire)", "JH Build", "JH Run/mois", "Notes"],
        "08_Risques": ["ID", "Risque", "Impact", "Probabilite", "Mitigation", "Responsable", "Statut"],
        "09_Arbitrages": ["ID", "Question", "Option", "Justification", "Statut", "Source"],
        "10_Sources": ["ID", "Titre", "Type", "Date", "URL", "Notes"],
        "11_Synthese": ["Indicateur", "Valeur", "Unite", "Notes"],
    }
    for name, headers in expected_headers.items():
        actual = ([formulas[name].cell(4, column).value for column in range(1, 7)]
                  if name == "06_Scenarios" and name in formulas.sheetnames
                  else ([cell.value for cell in formulas[name][4]] if name in formulas.sheetnames else headers))
        if actual != headers:
            report.errors.append(f"{name}: column contract mismatch")
    if "06_Scenarios" in formulas.sheetnames:
        ws = formulas["06_Scenarios"]
        covered = set()
        for table in ws.tables.values():
            left, top, right, bottom = range_boundaries(table.ref)
            if (left, top) == (1, 4):
                if right != 6:
                    report.errors.append("06_Scenarios: canonical table must be A:F")
                continue
            report.counts["annex_blocks"] += 1
            headers = [ws.cell(top, c).value for c in range(left, right + 1)]
            if any(not isinstance(h, str) or not h.strip() for h in headers) or len(headers) != len(set(headers)):
                report.errors.append(f"06_Scenarios: annex {table.name} has invalid headers")
            covered.update((r, c) for r in range(top, bottom + 1) for c in range(left, right + 1))
        if any(cell.value is not None and (cell.row, cell.column) not in covered
               for row in ws.iter_rows(min_row=4, min_col=7) for cell in row):
            report.errors.append("06_Scenarios: annex outside A:F requires a separately validated table")
    if all(name in formulas.sheetnames for name in specification["sheets"]):
        try:
            metadata = {row["Cle"]: row["Valeur"] for _, row in _rows(values, "00_Lisez-moi")}
            if metadata.get("effort_unit") != "JH":
                report.errors.append("Effort unit must be JH")
            currency = metadata["currency"]
            cost_headers = ["ID", "Fournisseur", "Service", "Environnement", "Scenario", "Unite", "Quantite",
                            f"Prix unitaire ({currency})", "Devise", "Multiplicateur mensuel", f"Cout mensuel ({currency})",
                            f"Cout annuel ({currency})", "Source", "Confiance", "Notes", "Statut prix"]
            if [cell.value for cell in formulas["05_Couts_cloud"][4]] != cost_headers:
                report.errors.append("05_Couts_cloud: column contract mismatch")
            _reconcile(formulas, values, specification, report, currency)
        except (KeyError, TypeError, ValueError, IndexError) as error:
            report.errors.append(f"Financial reconciliation unavailable: malformed or uncached data ({error})")
    if diagram_inventory_ids is not None:
        supplied = list(diagram_inventory_ids)
        report.counts["diagram_inventory_ids"] = len(supplied)
        inventory = {row["ID"] for _, row in _rows(values, "02_Inventaire")} if "02_Inventaire" in values.sheetnames else set()
        if len(supplied) != len(set(supplied)) or set(supplied) != inventory:
            report.errors.append(f"Diagram inventory mismatch: missing={sorted(inventory - set(supplied))}; unknown={sorted(set(supplied) - inventory)}")
    for reference, expected in (asserted_totals or {}).items():
        try:
            sheet, coordinate = reference.rsplit("!", 1)
            _compare(report, reference, values[sheet][coordinate].value, expected)
        except (ValueError, KeyError) as error:
            report.errors.append(f"Prose reconciliation target {reference} is invalid: {error}")
    report.warnings.append("Numerical prose and diagram correspondence are checked only for supplied assertions/IDs.")
    report.counts["errors"] = len(report.errors)
    report.ok = not report.errors
    return report


def _compare(report, location, actual, expected):
    report.counts["reconciliations"] += 1
    if isinstance(expected, (int, float)):
        match = isinstance(actual, (int, float)) and math.isfinite(actual) and math.isclose(actual, expected, rel_tol=1e-9, abs_tol=1e-7)
    else:
        match = actual == expected
    if not match:
        report.counts["reconciliation_errors"] += 1
        report.errors.append(f"{location}: reconciliation mismatch (actual {actual!r}, expected {expected!r})")


def _reconcile(formulas, values, contract, report, currency):
    assumptions = _assumptions(values)
    rate = assumptions["taux_journalier"]
    unknown_rate = rate in (None, "A chiffrer")
    report.finance_status = "blocked" if unknown_rate else "rate_provided_not_approved"
    if unknown_rate:
        report.warnings.append("Finance blocked: taux_journalier A chiffrer; dependent human costs and TCO unavailable. Technical capacity and JH remain valid.")
    elif isinstance(rate, bool) or not isinstance(rate, (int, float)) or not math.isfinite(rate) or rate < 0:
        report.errors.append("01_Hypotheses: invalid taux_journalier")
        return
    scenarios = contract["scenario_names"]
    global_scenario = assumptions["scenario_global"]
    if "comparison_scenario" in assumptions and assumptions["comparison_scenario"] not in scenarios:
        report.errors.append("01_Hypotheses: invalid comparison_scenario")
    if global_scenario not in scenarios:
        report.errors.append("01_Hypotheses: invalid global scenario")
        return
    tasks = _rows(values, "04_Taches_JH")
    effort = dict(zip(scenarios, ("JH bas", "JH probable", "JH haut")))
    totals = defaultdict(float)
    phases = defaultdict(float)
    all_scenarios = defaultdict(float)

    def derived(sheet, coordinate, expected):
        if formulas[sheet][coordinate].data_type != "f":
            report.errors.append(f"{sheet}!{coordinate}: derived value must remain a formula")
        _compare(report, f"{sheet}!{coordinate}", values[sheet][coordinate].value, expected)

    for index, task in tasks:
        selected = task["Scenario override"] or global_scenario
        if selected not in scenarios or task["Nature"] not in ("Build", "Run"):
            report.errors.append(f"04_Taches_JH!A{index}: invalid scenario or Nature")
            continue
        estimates = [task[column] for column in effort.values()]
        if any(not isinstance(item, (int, float)) or isinstance(item, bool) or not math.isfinite(item) or item < 0 for item in estimates) or estimates != sorted(estimates):
            report.errors.append(f"04_Taches_JH!L{index}: invalid JH estimates")
            continue
        selected_effort = task[effort[selected]]
        derived("04_Taches_JH", f"O{index}", selected)
        derived("04_Taches_JH", f"Q{index}", selected_effort)
        totals[task["Nature"]] += selected_effort
        phases[(task["Lot"], task["Nature"])] += selected_effort
        for scenario, column in effort.items():
            all_scenarios[(scenario, task["Nature"])] += task[column]
    cost_groups = defaultdict(float)
    provider_totals = defaultdict(float)
    for index, cost in _rows(values, "05_Couts_cloud"):
        if cost["Scenario"] not in scenarios or cost["Devise"] != currency:
            report.errors.append(f"05_Couts_cloud!A{index}: invalid scenario or currency")
        inputs = [cost["Quantite"], cost[f"Prix unitaire ({currency})"], cost["Multiplicateur mensuel"]]
        if any(not isinstance(item, (int, float)) or isinstance(item, bool) or not math.isfinite(item) or item < 0 for item in inputs):
            report.errors.append(f"05_Couts_cloud!G{index}: invalid cost inputs")
            continue
        monthly = math.prod(inputs)
        derived("05_Couts_cloud", f"K{index}", monthly)
        derived("05_Couts_cloud", f"L{index}", monthly * assumptions["mois_par_an"])
        group = (cost["Fournisseur"], cost["Environnement"], cost["Scenario"])
        cost_groups[group] += monthly
        provider_totals[(cost["Fournisseur"], cost["Scenario"])] += monthly
    report.cost_totals = [dict(provider=key[0], environment=key[1], scenario=key[2], monthly=value,
                               annual=value * assumptions["mois_par_an"], currency=currency)
                          for key, value in sorted(cost_groups.items())]
    for index, row in _rows(values, "06_Scenarios"):
        build, run = all_scenarios[(row["Scenario"], "Build")], all_scenarios[(row["Scenario"], "Run")]
        for column, amount in (("C", build + run), ("D", build), ("E", run)):
            derived("06_Scenarios", f"{column}{index}", amount)
    phase_names = set()
    for index, row in _rows(values, "07_Phasage"):
        phase_names.add(row["Phase"])
        build, run = phases[(row["Phase"], "Build")], phases[(row["Phase"], "Run")]
        for column, amount in (("D", build + run), ("E", build), ("F", run)):
            derived("07_Phasage", f"{column}{index}", amount)
    if {task["Lot"] for _, task in tasks} - phase_names:
        report.errors.append("07_Phasage: phases do not cover every task lot")
    expected_summary = {"Rapprochement technique (Build + Run/mois)": totals["Build"] + totals["Run"],
                        "JH Build": totals["Build"], "JH Run mensuel": totals["Run"],
                        "Cout Build": None if unknown_rate else totals["Build"] * rate,
                        "Cout Run mensuel": None if unknown_rate else totals["Run"] * rate,
                        "Statut finance": "Bloque - taux journalier A chiffrer" if unknown_rate else "Taux journalier renseigne - approbation non acquise"}
    providers = {provider for provider, _ in provider_totals}
    for provider in providers:
        for scenario in scenarios:
            expected_summary[f"Cloud mensuel | {provider} | {scenario}"] = provider_totals[(provider, scenario)]
    found = set()
    for index, row in _rows(values, "11_Synthese"):
        label = row["Indicateur"]
        if label not in expected_summary:
            report.errors.append(f"11_Synthese!A{index}: unknown summary indicator")
            continue
        if label in found:
            report.errors.append(f"11_Synthese!A{index}: duplicate summary indicator")
        found.add(label)
        derived("11_Synthese", f"B{index}", expected_summary[label])
        if label.startswith("Rapprochement technique") and (row["Unite"] != "Unites mixtes" or "utilisation budgetaire interdite" not in str(row["Notes"]).lower()):
            report.errors.append(f"11_Synthese!D{index}: mixed-unit total must prohibit budget use")
    if set(expected_summary) - found:
        report.errors.append(f"11_Synthese: missing summary indicators {sorted(set(expected_summary) - found)}")
