"""Exercise the shipped synthetic workbook, including cached domain calculations.

Missing caches, a merged provider budget, dropped flows, double-counted tasks,
or lost capacity/FinOps decisions must fail these consumer-facing checks.
"""

from collections import Counter
import json
from pathlib import Path
import subprocess
import sys

from openpyxl import load_workbook
import pytest

from tooling.anonymisation import scan_tree
from tooling.model import load_model
from tooling.recalculate import FormulaFailure, _run_excel_powershell
from tooling.workbook_validator import validate_workbook

ROOT = Path(__file__).parents[1]
EXAMPLES = ROOT / "examples"


@pytest.fixture(scope="module")
def example_report():
    path = EXAMPLES / "synthetic-cloud-estimate.xlsx"
    for name in ("synthetic-cloud-estimate.xlsx", "synthetic-cloud-brief.json",
                 "synthetic-architecture.json", "synthetic-assertions.json"):
        assert (EXAMPLES / name).is_file(), f"Generate the complete public example: missing {name}"
    model = load_model(EXAMPLES / "synthetic-cloud-brief.json")
    formulas = load_workbook(path)
    values = load_workbook(path, data_only=True)
    architecture = json.loads((EXAMPLES / "synthetic-architecture.json").read_text(encoding="utf-8"))
    inventory = [row.id if hasattr(row, "id") else row["id"] for row in model.inventory]
    assertions = json.loads((EXAMPLES / "synthetic-assertions.json").read_text(encoding="utf-8"))
    report = validate_workbook(path, diagram_inventory_ids=inventory, asserted_totals=assertions)
    return model, formulas, values, architecture, report


def test_public_example_meets_suite_success_criteria(example_report):
    assert example_report is not None, "Generate the public synthetic XLSX acceptance example"
    model, formulas, values, architecture, report = example_report
    assert report.ok, report.errors
    assert report.counts["formula_errors"] == 0
    assert report.counts["missing_cached_values"] == 0
    assert report.counts["formula_semantic_errors"] == 0
    assert report.counts["reconciliation_errors"] == 0
    assert report.counts["formulas"] >= 100
    assert report.counts["reconciliations"] >= 100
    assert len(model.scenarios) == 3
    providers = {row.provider for row in model.cloud_costs}
    assert len(providers) == 3
    groups = Counter((row.provider, row.scenario) for row in model.cloud_costs)
    assert len(groups) == 9 and all(count == 3 for count in groups.values())
    assert len({row.id for row in model.tasks}) == len(model.tasks) == 10
    inventory_ids = {row["id"] for row in model.inventory}
    assert len(inventory_ids) == 4
    manifest = next(json.loads(row[1].value) for row in formulas["03_Architecture"]
                    if row[0].value == "JSON")
    assert set(manifest["node_ids"]) == inventory_ids
    assert set(manifest["edge_ids"]) == {row["id"] for row in architecture["edges"]}
    assert len(manifest["edge_ids"]) == 1
    assert manifest["icon_provenance"] == []
    assert len(formulas["03_Architecture"]._images) == 1
    scan = scan_tree(ROOT)
    assert scan.clean, scan.summary()
    assert scan.match_count == 0 and scan.scanned_files > 200


def test_domain_calculations_and_separate_build_run_reconcile(example_report):
    assert example_report is not None, "Generate the public synthetic XLSX acceptance example"
    model, formulas, values, architecture, report = example_report
    assumptions = {row[0].value: row[1] for row in values["01_Hypotheses"] if row[0].value}
    expected = {
        "schedule.foundation_overload": 3,
        "capacity.production.Standard.high.overload_ratio": 2,
        "contradiction.cloud_delta": 45,
        "contradiction.build_delta": 2,
        "contradiction.run_delta": 0.5,
        "finops.savings_low": 22.5,
        "finops.savings_high": 45,
        "finops.savings_selected": 45,
        "tco.ProviderA.Standard": 41906,
    }
    for key, value in expected.items():
        cell = assumptions[key]
        assert cell.value == pytest.approx(value), key
        assert formulas["01_Hypotheses"][cell.coordinate].data_type == "f", key
    assert values["11_Synthese"]["B6"].value == 19
    assert values["11_Synthese"]["B7"].value == 1
    assert values["11_Synthese"]["C6"].value == "JH"
    assert values["11_Synthese"]["C7"].value == "JH/mois"
    assert values["11_Synthese"]["B8"].value == 12350
    assert values["11_Synthese"]["B9"].value == 650
    assert values["11_Synthese"]["B11"].value == 171
    assert values["06_Scenarios"]["D5"].value == 9.5
    assert values["06_Scenarios"]["D6"].value == 17
    assert values["06_Scenarios"]["D7"].value == 30
    assert {row["id"] for row in model.decisions} >= {"DEC-CONFLICT", "DEC-FINOPS", "DEC-ARCH", "DEC-SCHEDULE"}
    assert any(row.domain == "finops" and row.nature == "Build" for row in model.tasks)
    assert any(row.domain == "finops" and row.nature == "Run" for row in model.tasks)


def test_regeneration_runs_real_shared_toolchain(tmp_path):
    script = EXAMPLES / "generate_example.py"
    assert script.is_file(), "Missing executable public example generator"
    result = subprocess.run([sys.executable, str(script), "--output-dir", str(tmp_path)],
                            capture_output=True, text=True, timeout=120)
    assert result.returncode == 0, result.stdout + result.stderr
    evidence = json.loads(result.stdout)
    assert evidence["validation"]["ok"]
    assert evidence["recalculation"]["formula_count"] >= 100
    assert evidence["recalculation"]["cached_values_count"] == evidence["recalculation"]["formula_count"]
    assert evidence["anonymisation_matches"] == 0
    assert not list(tmp_path.glob("estimate-recalc-*")), "Native recalculation left temporary files"
    path = tmp_path / "synthetic-cloud-estimate.xlsx"
    with path.open("rb") as stream:
        assert load_workbook(stream, data_only=True)["11_Synthese"]["B6"].value == 19
    # On Windows these operations fail if any verification reader retains a handle.
    replacement = tmp_path / "replacement.xlsx"
    path.replace(replacement)
    replacement.unlink()
    assert not path.exists() and not replacement.exists()


@pytest.mark.skipif(sys.platform != "win32", reason="Windows PowerShell COM bridge")
def test_windows_com_bridge_rejects_exit_zero_before_completion():
    _run_excel_powershell("Write-Output 'control'")
    with pytest.raises(FormulaFailure, match="PowerShell Excel COM failed"):
        _run_excel_powershell("exit 0")
