import json
from dataclasses import FrozenInstanceError
from pathlib import Path

import pytest
from openpyxl import load_workbook

from tooling.model import load_model
from tooling.workbook_builder import build_workbook

ROOT = Path(__file__).parents[1]
FIXTURE = ROOT / "tests/fixtures/synthetic-brief.json"


def test_builder_emits_contract_and_formulas(tmp_path):
    model = load_model(FIXTURE)
    wb = load_workbook(build_workbook(model, tmp_path / "estimate.xlsx"))
    contract = json.loads((ROOT / "tooling/workbook-contract.json").read_text(encoding="utf-8"))
    assert wb.sheetnames == contract["sheets"]
    ws = wb["04_Taches_JH"]
    headers = [cell.value for cell in ws[4]]
    assert headers == contract["task_columns"]
    assert ws.cell(5, headers.index("JH retenu") + 1).value.startswith("=")
    assert len(model.tasks) == 6
    assert {row["environment"] for row in model.inventory} == {"Development", "Staging", "Production"}
    assert {cost.provider for cost in model.cloud_costs} == {"Provider A", "Provider B", "Provider C"}


def test_builder_uses_required_financial_colours(tmp_path):
    wb = load_workbook(build_workbook(load_model(FIXTURE), tmp_path / "estimate.xlsx"))
    assert wb["01_Hypotheses"]["B5"].font.color.rgb.endswith("0000FF")
    assert wb["01_Hypotheses"]["B6"].fill.fgColor.rgb.endswith("FFFF00")
    assert wb["04_Taches_JH"]["Q5"].font.color.rgb.endswith("000000")
    assert wb["04_Taches_JH"]["O5"].font.color.rgb.endswith("008000")
    assert wb["11_Synthese"]["B5"].font.color.rgb.endswith("008000")
    assert wb["05_Couts_cloud"]["L5"].font.color.rgb.endswith("008000")


def test_mixed_effort_totals_prohibit_budget_use(tmp_path):
    wb = load_workbook(build_workbook(load_model(FIXTURE), tmp_path / "estimate.xlsx"))
    for name, total_column in (("06_Scenarios", "C4"), ("07_Phasage", "D4")):
        ws = wb[name]
        assert "rapprochement technique" in ws[total_column].value.lower()
        assert "utilisation budgetaire interdite" in ws["A2"].value.lower()
        headers = [cell.value for cell in ws[4]]
        assert "JH Build" in headers
        assert "JH Run/mois" in headers
    summary = wb["11_Synthese"]
    assert "rapprochement technique" in summary["A5"].value.lower()
    assert summary["C5"].value == "Unites mixtes"
    assert "utilisation budgetaire interdite" in summary["D5"].value.lower()


def test_scenario_selection_and_override_are_dynamic(tmp_path):
    wb = load_workbook(build_workbook(load_model(FIXTURE), tmp_path / "estimate.xlsx"))
    ws = wb["04_Taches_JH"]
    assert ws["O5"].value == '=IF(P5="",scenario_global,P5)'
    assert ws["P6"].value == "Resilient"
    assert ws["Q5"].value == '=IF(O5="Economique",L5,IF(O5="Standard",M5,N5))'
    assert list(wb.defined_names["scenario_global"].destinations) == [("01_Hypotheses", "$B$5")]
    assert any("Economique,Standard,Resilient" in validation.formula1 for validation in ws.data_validations.dataValidation)


def test_all_sheets_have_titles_headers_and_non_overlapping_tables(tmp_path):
    wb = load_workbook(build_workbook(load_model(FIXTURE), tmp_path / "estimate.xlsx"))
    names = []
    for ws in wb:
        assert ws["A1"].value and ws["A2"].value and ws["A4"].value
        assert ws.freeze_panes == "A5"
        assert len(ws.tables) == 1
        assert ws.auto_filter.ref is None
        table = next(iter(ws.tables.values()))
        assert table.ref.startswith("A4:")
        names.append(table.name)
    assert len(set(names)) == 12
    assert wb.calculation.fullCalcOnLoad


def test_costs_and_scenario_and_phase_totals_are_formulas(tmp_path):
    wb = load_workbook(build_workbook(load_model(FIXTURE), tmp_path / "estimate.xlsx"))
    costs = wb["05_Couts_cloud"]
    assert costs["K5"].value == "=G5*H5*J5"
    assert costs["L5"].value == "=K5*mois_par_an"
    assert wb["06_Scenarios"]["C5"].value.startswith("=SUM(")
    assert wb["07_Phasage"]["D5"].value.startswith("=SUMIF(")
    assert "Taches[JH retenu]" in wb["11_Synthese"]["B5"].value
    assert "SUMIF" in wb["11_Synthese"]["B8"].value


def test_named_assumption_coordinates_follow_inserted_rows(tmp_path):
    data = json.loads(FIXTURE.read_text(encoding="utf-8"))
    data["assumptions"].insert(0, {"key": "extra", "value": "Synthetic", "unit": "", "status": "Confirme", "source": "SYN-01", "notes": ""})
    data["tasks"].append(dict(data["tasks"][0], id="TASK-07"))
    path = tmp_path / "brief.json"
    path.write_text(json.dumps(data), encoding="utf-8")
    wb = load_workbook(build_workbook(load_model(path), tmp_path / "estimate.xlsx"))
    assert list(wb.defined_names["scenario_global"].destinations) == [("01_Hypotheses", "$B$6")]
    assert wb["04_Taches_JH"].tables["Taches"].ref == "A4:V11"
    assert wb["11_Synthese"]["B5"].value == "=SUM(Taches[JH retenu])"


@pytest.mark.parametrize("section,field,value", [("tasks", "jh_low", -1), ("tasks", "jh_likely", float("nan")), ("tasks", "jh_high", True), ("tasks", "scenario_override", "Unknown"), ("cloud_costs", "quantity", -1), ("scenarios", "name", "Unknown")])
def test_model_rejects_invalid_values(tmp_path, section, field, value):
    data = json.loads(FIXTURE.read_text(encoding="utf-8"))
    data[section][0][field] = value
    path = tmp_path / "invalid.json"
    path.write_text(json.dumps(data), encoding="utf-8")
    with pytest.raises(ValueError, match=field):
        load_model(path)


@pytest.mark.parametrize("missing", ["metadata", "tasks", "sources"])
def test_model_rejects_missing_required_sections(tmp_path, missing):
    data = json.loads(FIXTURE.read_text(encoding="utf-8"))
    del data[missing]
    path = tmp_path / "invalid.json"
    path.write_text(json.dumps(data), encoding="utf-8")
    with pytest.raises(ValueError, match=missing):
        load_model(path)


def test_model_validates_task_ordering_ids_and_dependencies(tmp_path):
    data = json.loads(FIXTURE.read_text(encoding="utf-8"))
    for field, value, message in [("jh_low", 99, "jh_low"), ("id", "TASK-02", "duplicate"), ("dependencies", ["MISSING"], "dependencies")]:
        changed = json.loads(json.dumps(data))
        changed["tasks"][0][field] = value
        path = tmp_path / "invalid.json"
        path.write_text(json.dumps(changed), encoding="utf-8")
        with pytest.raises(ValueError, match=message):
            load_model(path)


def test_model_is_frozen_and_required_task_fields_are_enforced(tmp_path):
    model = load_model(FIXTURE)
    with pytest.raises(FrozenInstanceError):
        model.tasks[0].jh_low = 0
    data = json.loads(FIXTURE.read_text(encoding="utf-8"))
    del data["tasks"][0]["deliverable"]
    path = tmp_path / "invalid.json"
    path.write_text(json.dumps(data), encoding="utf-8")
    with pytest.raises(ValueError, match="deliverable"):
        load_model(path)


def test_empty_optional_sections_still_emit_tables(tmp_path):
    data = json.loads(FIXTURE.read_text(encoding="utf-8"))
    for section in ("inventory", "architecture", "cloud_costs", "phasing", "risks", "decisions", "sources"):
        data[section] = []
    path = tmp_path / "empty.json"
    path.write_text(json.dumps(data), encoding="utf-8")
    wb = load_workbook(build_workbook(load_model(path), tmp_path / "estimate.xlsx"))
    assert len(wb.sheetnames) == 12
    assert wb["05_Couts_cloud"].tables["CoutsCloud"].ref == "A4:P5"


def test_input_text_cannot_become_excel_formula(tmp_path):
    data = json.loads(FIXTURE.read_text(encoding="utf-8"))
    data["tasks"][0]["notes"] = "=1+1"
    path = tmp_path / "brief.json"
    path.write_text(json.dumps(data), encoding="utf-8")
    wb = load_workbook(build_workbook(load_model(path), tmp_path / "estimate.xlsx"))
    assert wb["04_Taches_JH"]["V5"].data_type == "s"


@pytest.mark.parametrize("field,value", [("id", 12), ("deliverable", None), ("nature", "Unknown")])
def test_model_rejects_invalid_typed_task_fields(tmp_path, field, value):
    data = json.loads(FIXTURE.read_text(encoding="utf-8"))
    data["tasks"][0][field] = value
    path = tmp_path / "invalid.json"
    path.write_text(json.dumps(data), encoding="utf-8")
    with pytest.raises(ValueError, match=field):
        load_model(path)


def test_model_rejects_dependency_cycles(tmp_path):
    data = json.loads(FIXTURE.read_text(encoding="utf-8"))
    data["tasks"][0]["dependencies"] = ["TASK-02"]
    path = tmp_path / "invalid.json"
    path.write_text(json.dumps(data), encoding="utf-8")
    with pytest.raises(ValueError, match="dependencies.*cycle"):
        load_model(path)


def test_model_accepts_zero_effort_and_zero_cost(tmp_path):
    data = json.loads(FIXTURE.read_text(encoding="utf-8"))
    data["tasks"][0].update(jh_low=0, jh_likely=0, jh_high=0)
    data["cloud_costs"][0].update(quantity=0, unit_price=0)
    path = tmp_path / "zero.json"
    path.write_text(json.dumps(data), encoding="utf-8")
    model = load_model(path)
    wb = load_workbook(build_workbook(model, tmp_path / "zero.xlsx"))
    assert wb["04_Taches_JH"]["L5"].value == 0
    assert wb["05_Couts_cloud"]["G5"].value == 0
