"""Observable diagrams, integrity refusal, and workbook integration."""

from copy import deepcopy
from hashlib import sha256
import json
from pathlib import Path
import re
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from PIL import Image, ImageFont
import pytest

from tooling.diagram_renderer import render_diagram
from tooling.model import load_model
from tooling.workbook_builder import build_workbook

ROOT = Path(__file__).parents[1]
NS = {"s": "http://www.w3.org/2000/svg"}


@pytest.fixture
def architecture():
    return {
        "nodes": [
            {"id": "COMP-03", "component": "Database", "environment": "Production", "zone": "Private"},
            {"id": "COMP-01", "component": "Web <service>", "environment": "Development", "zone": "Public"},
            {"id": "COMP-02", "component": "API", "environment": "Production", "zone": "Private"},
        ],
        "edges": [{"id": "FLOW-01", "source": "COMP-02", "target": "COMP-03", "label": "TLS"}],
    }


def test_diagram_contains_inventory_nodes_and_directed_flows(tmp_path, architecture):
    result = render_diagram(architecture, tmp_path / "a.svg", tmp_path / "a.png")
    assert result.node_ids == ["COMP-01", "COMP-02", "COMP-03"]
    assert result.edge_ids == ["FLOW-01"]
    svg = ET.parse(tmp_path / "a.svg")
    assert {node.attrib["data-node-id"] for node in svg.findall(".//s:g[@data-node-id]", NS)} == set(result.node_ids)
    edge = svg.find(".//s:path[@data-edge-id]", NS)
    assert edge.attrib["data-edge-id"] == "FLOW-01"
    assert edge.attrib["data-source"] == "COMP-02"
    assert edge.attrib["data-target"] == "COMP-03"
    assert edge.attrib["marker-end"] == "url(#arrow)"
    assert {text.text for text in svg.findall(".//s:text", NS)} >= {"Development", "Production", "Private", "Public", "TLS", "Web <service>"}
    assert (tmp_path / "a.svg").read_text(encoding="utf-8").startswith("<svg")
    with Image.open(tmp_path / "a.png") as png:
        assert png.format == "PNG"
        assert png.width > 100 and png.height > 100
        assert len(png.getcolors(png.width * png.height)) > 2
    assert result.output_hashes == {"svg": sha256((tmp_path / "a.svg").read_bytes()).hexdigest(), "png": sha256((tmp_path / "a.png").read_bytes()).hexdigest()}


def test_reordered_input_produces_identical_outputs(tmp_path, architecture):
    first = render_diagram(architecture, tmp_path / "a.svg", tmp_path / "a.png")
    changed = deepcopy(architecture)
    changed["nodes"].reverse()
    second = render_diagram(changed, tmp_path / "b.svg", tmp_path / "b.png")
    assert second.output_hashes == first.output_hashes
    assert second.to_dict() == first.to_dict()


def test_flow_routes_around_intervening_nodes(tmp_path):
    architecture = {"nodes": [{"id": key, "environment": "Production", "zone": "Private"} for key in ("A", "B", "C")],
                    "edges": [{"id": "C-to-A", "source": "C", "target": "A"}]}
    render_diagram(architecture, tmp_path / "a.svg", tmp_path / "a.png")
    svg = ET.parse(tmp_path / "a.svg")
    middle = svg.find(".//s:g[@data-node-id='B']/s:rect", NS).attrib
    x, y, w, h = (float(middle[key]) for key in ("x", "y", "width", "height"))
    coordinates = [float(value) for value in re.findall(r"-?\d+(?:\.\d+)?", svg.find(".//s:path[@data-edge-id]", NS).attrib["d"])]
    points = list(zip(coordinates[::2], coordinates[1::2]))
    assert len(points) >= 2
    for (ax, ay), (bx, by) in zip(points, points[1:]):
        for step in range(21):
            sx, sy = ax + (bx - ax) * step / 20, ay + (by - ay) * step / 20
            assert not (x < sx < x + w and y < sy < y + h), "flow crosses unrelated node"


def assert_paths_avoid_node_interiors(svg):
    rectangles = [node.attrib for node in svg.findall(".//s:g[@data-node-id]/s:rect", NS)]
    paths = svg.findall(".//s:path[@data-edge-id]", NS)
    assert rectangles and paths
    for path in paths:
        coordinates = [float(value) for value in re.findall(r"-?\d+(?:\.\d+)?", path.attrib["d"])]
        points = list(zip(coordinates[::2], coordinates[1::2]))
        assert len(points) >= 2
        for (ax, ay), (bx, by) in zip(points, points[1:]):
            assert ax == bx or ay == by, "expected orthogonal route"
            for rectangle in rectangles:
                x, y, w, h = (float(rectangle[key]) for key in ("x", "y", "width", "height"))
                if ax == bx:
                    intersects = x < ax < x + w and max(min(ay, by), y) < min(max(ay, by), y + h)
                else:
                    intersects = y < ay < y + h and max(min(ax, bx), x) < min(max(ax, bx), x + w)
                assert not intersects, f"{path.attrib['data-edge-id']} crosses node rectangle {rectangle}"


@pytest.mark.parametrize("reverse", [False, True])
def test_cross_environment_routes_avoid_all_nodes(tmp_path, reverse):
    nodes = [{"id": key, "environment": "A", "zone": "One"} for key in ("A-1", "A-2", "A-3")]
    nodes += [{"id": "B-1", "environment": "B", "zone": "One"}, {"id": "B-2", "environment": "B", "zone": "Two"}]
    source, target = ("A-1", "B-2") if reverse else ("B-2", "A-1")
    architecture = {"nodes": nodes, "edges": [{"id": "cross-environment", "source": source, "target": target}]}
    first = render_diagram(architecture, tmp_path / "a.svg", tmp_path / "a.png")
    assert_paths_avoid_node_interiors(ET.parse(tmp_path / "a.svg"))
    architecture["nodes"].reverse()
    second = render_diagram(architecture, tmp_path / "b.svg", tmp_path / "b.png")
    assert first.output_hashes == second.output_hashes


def test_multirow_routes_avoid_every_node(tmp_path):
    nodes = [{"id": f"N-{index}", "environment": "A", "zone": "One"} for index in range(8)]
    edges = [{"id": f"FLOW-{index}", "source": f"N-{index}", "target": "N-0"} for index in range(1, 8)]
    render_diagram({"nodes": nodes, "edges": edges}, tmp_path / "a.svg", tmp_path / "a.png")
    assert_paths_avoid_node_interiors(ET.parse(tmp_path / "a.svg"))


@pytest.mark.parametrize("change,message", [
    (lambda a: a["nodes"].append(dict(a["nodes"][0])), "duplicate node"),
    (lambda a: a["edges"].append(dict(a["edges"][0])), "duplicate edge"),
    (lambda a: a["edges"][0].update(target="UNKNOWN"), "unknown node"),
    (lambda a: a["nodes"][0].update(id=""), "node.*id"),
    (lambda a: a["nodes"][0].update(component="invalid\ufffe"), "invalid XML text"),
])
def test_invalid_graph_is_refused_before_outputs(tmp_path, architecture, change, message):
    change(architecture)
    with pytest.raises(ValueError, match=message):
        render_diagram(architecture, tmp_path / "a.svg", tmp_path / "a.png")
    assert not (tmp_path / "a.svg").exists()
    assert not (tmp_path / "a.png").exists()


def make_icon_pack(tmp_path):
    """Tiny generic test image; never a generated proprietary service logo."""
    pack = tmp_path / "icons"
    pack.mkdir()
    asset = pack / "symbol.png"
    Image.new("RGB", (24, 24), "#ba2020").save(asset)
    manifest = {
        "provider": "AWS", "official_source_url": "https://aws.amazon.com/architecture/icons/",
        "trademark_notice": "Test provenance notice; fixture contains no provider logo.",
        "assets": [{"path": "symbol.png", "sha256": sha256(asset.read_bytes()).hexdigest()}],
    }
    (pack / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    return pack, manifest


@pytest.mark.parametrize("defect", ["missing-manifest", "missing-notice", "bad-hash", "unofficial-source", "host-spoof", "path-escape", "missing-asset", "unlisted-asset", "invalid-json", "invalid-image"])
def test_unverified_icon_packs_fall_back_to_neutral(tmp_path, architecture, defect):
    pack, manifest = make_icon_pack(tmp_path)
    architecture["nodes"][0].update(provider="AWS", icon="symbol.png")
    if defect == "missing-notice":
        del manifest["trademark_notice"]
    elif defect == "bad-hash":
        manifest["assets"][0]["sha256"] = "0" * 64
    elif defect == "unofficial-source":
        manifest["official_source_url"] = "https://example.com/aws/icons/"
    elif defect == "host-spoof":
        manifest["official_source_url"] = "https://aws.amazon.com.example.com/architecture/icons/"
    elif defect == "path-escape":
        manifest["assets"][0]["path"] = "../symbol.png"
    elif defect == "missing-asset":
        manifest["assets"][0]["path"] = "missing.png"
    elif defect == "unlisted-asset":
        Image.new("RGB", (2, 2)).save(pack / "unlisted.png")
    elif defect == "invalid-image":
        (pack / "symbol.png").write_bytes(b"not an image")
        manifest["assets"][0]["sha256"] = sha256(b"not an image").hexdigest()
    if defect == "missing-manifest":
        (pack / "manifest.json").unlink()
    else:
        (pack / "manifest.json").write_text("{" if defect == "invalid-json" else json.dumps(manifest), encoding="utf-8")
    result = render_diagram(architecture, tmp_path / "a.svg", tmp_path / "a.png", pack)
    assert result.icon_sources == ["provider-neutral"]
    assert result.icon_provenance == []
    assert not ET.parse(tmp_path / "a.svg").findall(".//s:image", NS)


def test_verified_icon_is_embedded_with_provenance(tmp_path, architecture):
    pack, manifest = make_icon_pack(tmp_path)
    architecture["nodes"][0].update(provider="AWS", icon="symbol.png")
    result = render_diagram(architecture, tmp_path / "a.svg", tmp_path / "a.png", pack)
    assert result.icon_sources == ["provider-neutral", manifest["official_source_url"]]
    assert result.icon_provenance == [{**manifest["assets"][0], "provider": "AWS", "official_source_url": manifest["official_source_url"], "licence_or_trademark_notice": manifest["trademark_notice"]}]
    image = ET.parse(tmp_path / "a.svg").find(".//s:image", NS)
    assert image.attrib["href"].startswith("data:image/png;base64,")
    with Image.open(tmp_path / "a.png") as png:
        assert any(colour == (186, 32, 32) for _, colour in png.getcolors(png.width * png.height))
    assert str(pack) not in json.dumps(result.to_dict())


def test_provider_mismatch_uses_neutral_symbol(tmp_path, architecture):
    pack, _ = make_icon_pack(tmp_path)
    architecture["nodes"][0].update(provider="Azure", icon="symbol.png")
    result = render_diagram(architecture, tmp_path / "a.svg", tmp_path / "a.png", pack)
    assert result.icon_sources == ["provider-neutral"]
    assert result.icon_provenance == []


@pytest.mark.parametrize("invalid_character", ["\ufffe", "\x00"])
@pytest.mark.parametrize("notice_key", ["licence", "license", "trademark_notice"])
def test_invalid_rights_notice_falls_back_in_svg_and_workbook(tmp_path, architecture, invalid_character, notice_key):
    pack, manifest = make_icon_pack(tmp_path)
    manifest[notice_key] = "Test rights notice" + invalid_character
    (pack / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    architecture["nodes"][0].update(provider="AWS", icon="symbol.png")
    result = render_diagram(architecture, tmp_path / "a.svg", tmp_path / "a.png", pack)
    assert result.icon_sources == ["provider-neutral"]
    assert result.icon_provenance == []
    assert not ET.parse(tmp_path / "a.svg").findall(".//s:image", NS)
    model = load_model(ROOT / "tests/fixtures/synthetic-brief.json")
    wb = load_workbook(build_workbook(model, tmp_path / "estimate.xlsx", architecture=architecture, icon_pack=pack))
    ws = wb["03_Architecture"]
    marker = next(cell for row in ws for cell in row if cell.value == "Diagram manifest")
    stored = json.loads(ws.cell(marker.row + 1, 2).value)
    assert stored["icon_sources"] == ["provider-neutral"]
    assert stored["icon_provenance"] == []


def test_font_loading_supports_pillow_without_size_parameter(tmp_path, architecture, monkeypatch):
    original = ImageFont.load_default

    def legacy_default():
        return original()

    monkeypatch.setattr(ImageFont, "load_default", legacy_default)
    result = render_diagram(architecture, tmp_path / "a.svg", tmp_path / "a.png")
    assert result.node_ids == ["COMP-01", "COMP-02", "COMP-03"]
    with Image.open(tmp_path / "a.png") as png:
        assert png.format == "PNG" and png.width > 100


def test_workbook_embeds_image_manifest_inventory_and_flow_ranges(tmp_path, architecture):
    model = load_model(ROOT / "tests/fixtures/synthetic-brief.json")
    output = build_workbook(model, tmp_path / "estimate.xlsx", architecture=architecture)
    wb = load_workbook(output)
    ws = wb["03_Architecture"]
    assert len(ws._images) == 1
    cells = [cell.value for row in ws for cell in row]
    assert "Diagram nodes" in cells and "Diagram flows" in cells
    assert {"COMP-01", "COMP-02", "COMP-03", "FLOW-01"}.issubset(cells)
    assert "provider-neutral" in cells
    marker = next(cell for row in ws for cell in row if cell.value == "Diagram manifest")
    manifest = json.loads(ws.cell(marker.row + 1, 2).value)
    assert set(manifest["node_ids"]) == {row["id"] for row in model.inventory}
    assert manifest["edge_ids"] == ["FLOW-01"]
    with ZipFile(output) as archive:
        assert sha256(archive.read("xl/media/image1.png")).hexdigest() == manifest["output_hashes"]["png"]
    assert "icons" not in json.dumps(manifest)


def test_workbook_default_diagram_uses_inventory_ids(tmp_path):
    model = load_model(ROOT / "tests/fixtures/synthetic-brief.json")
    wb = load_workbook(build_workbook(model, tmp_path / "estimate.xlsx"))
    ws = wb["03_Architecture"]
    marker = next(cell for row in ws for cell in row if cell.value == "Diagram manifest")
    manifest = json.loads(ws.cell(marker.row + 1, 2).value)
    assert manifest["node_ids"] == ["COMP-01", "COMP-02", "COMP-03"]
    assert manifest["edge_ids"] == []
    assert len(ws.tables) == 1


def test_workbook_refuses_diagram_inventory_mismatch(tmp_path, architecture):
    model = load_model(ROOT / "tests/fixtures/synthetic-brief.json")
    architecture["nodes"].pop(1)  # Remove a node unconnected to the retained flow.
    with pytest.raises(ValueError, match="inventory"):
        build_workbook(model, tmp_path / "estimate.xlsx", architecture=architecture)
    assert not (tmp_path / "estimate.xlsx").exists()
