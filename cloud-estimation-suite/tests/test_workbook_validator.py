"""Negative controls: each corruption must change the delivery verdict."""

import json
from pathlib import Path
import sys
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from tooling.model import load_model
from tooling.workbook_builder import build_workbook
from tooling.recalculate import recalculate
from tooling.workbook_validator import validate_workbook

ROOT = Path(__file__).parents[1]
FIXTURE = ROOT / "tests/fixtures/synthetic-brief.json"


@pytest.fixture
def valid_workbook(tmp_path, monkeypatch):
    import tooling.recalculate as module
    monkeypatch.setattr(module, "_select_excel_backend", lambda: None)
    monkeypatch.setattr(module, "_find_libreoffice", lambda: None)
    path = build_workbook(load_model(FIXTURE), tmp_path / "estimate.xlsx")
    result = recalculate(path)
    assert result.ok, result.errors
    return path


def mutate(path, sheet, cell, value):
    wb = load_workbook(path)
    wb[sheet][cell] = value
    wb.save(path)


def test_validator_rejects_formula_errors(tmp_path):
    path = tmp_path / "broken.xlsx"
    wb = Workbook()
    wb.active["A1"] = "=#REF!"
    wb.save(path)
    report = validate_workbook(path)
    assert not report.ok
    assert any("#REF!" in error and "A1" in error for error in report.errors)


def test_roundtrip_populates_caches_preserves_formulas_and_reports_zero_errors(valid_workbook):
    report = validate_workbook(valid_workbook)
    assert report.ok, report.errors
    assert report.counts["errors"] == 0
    assert report.counts["missing_cached_values"] == 0
    assert report.counts["formula_errors"] == 0
    assert report.counts["reconciliation_errors"] == 0
    assert report.counts["formulas"] > 40
    assert json.loads(json.dumps(report.to_dict()))["ok"]
    values = load_workbook(valid_workbook, data_only=True)
    assert values["11_Synthese"]["B6"].value == 14
    assert values["11_Synthese"]["B7"].value == 1
    assert values["11_Synthese"]["B8"].value == 9100
    assert values["05_Couts_cloud"]["L5"].value == 1080
    assert load_workbook(valid_workbook)["11_Synthese"]["B6"].data_type == "f"


def test_validator_rejects_uncached_delivery(tmp_path):
    path = build_workbook(load_model(FIXTURE), tmp_path / "uncached.xlsx")
    report = validate_workbook(path)
    assert not report.ok
    assert report.counts["missing_cached_values"] > 0


def test_validator_rejects_hardcoded_prose_total(valid_workbook):
    report = validate_workbook(valid_workbook, asserted_totals={"11_Synthese!B6": 999})
    assert not report.ok
    assert any("reconciliation" in error.lower() for error in report.errors)


@pytest.mark.parametrize("sheet,cell,formula", [
    ("11_Synthese", "B6", "=1+1"),
    ("04_Taches_JH", "Q5", "=99"),
    ("05_Couts_cloud", "K5", "=1"),
    ("05_Couts_cloud", "L5", "=1"),
    ("06_Scenarios", "D5", "=1"),
    ("07_Phasage", "E5", "=1"),
    ("11_Synthese", "B11", "=1"),
])
def test_validator_reconciles_financial_totals(valid_workbook, sheet, cell, formula):
    mutate(valid_workbook, sheet, cell, formula)
    assert recalculate(valid_workbook).ok
    report = validate_workbook(valid_workbook)
    assert not report.ok
    assert any(f"{sheet}!{cell}" in error for error in report.errors)


def test_validator_rejects_hardcoded_derived_value(valid_workbook):
    mutate(valid_workbook, "11_Synthese", "B6", 14)
    assert recalculate(valid_workbook).ok
    report = validate_workbook(valid_workbook)
    assert not report.ok
    assert any("formula" in error.lower() and "B6" in error for error in report.errors)


def test_validator_rejects_filter_overlap(valid_workbook):
    wb = load_workbook(valid_workbook)
    ws = wb["04_Taches_JH"]
    ws.auto_filter.ref = "A4:C7"
    wb.save(valid_workbook)
    report = validate_workbook(valid_workbook)
    assert not report.ok
    assert report.counts["table_filter_overlaps"] == 1


@pytest.mark.parametrize("corruption", ["sheet", "column", "diagram", "future", "private"])
def test_validator_rejects_structural_and_publication_defects(valid_workbook, corruption):
    wb = load_workbook(valid_workbook)
    kwargs = {}
    if corruption == "sheet":
        del wb["08_Risques"]
    elif corruption == "column":
        wb["04_Taches_JH"]["L4"] = "Hours"
    elif corruption == "diagram":
        kwargs["diagram_inventory_ids"] = ["COMP-MISSING"]
    elif corruption == "future":
        wb["11_Synthese"]["B6"] = '=IFNA(A1,0)'
    else:
        wb["00_Lisez-moi"]["B5"] = "tenant_id=private-example"
    if corruption != "diagram":
        wb.save(valid_workbook)
    report = validate_workbook(valid_workbook, **kwargs)
    assert not report.ok
    assert any(corruption in error.lower() for error in report.errors)


def test_python_fallback_calibrates_nested_and_multiple_criteria_controls(tmp_path, monkeypatch):
    import tooling.recalculate as module
    monkeypatch.setattr(module, "_select_excel_backend", lambda: None)
    monkeypatch.setattr(module, "_find_libreoffice", lambda: None)
    path = build_workbook(load_model(FIXTURE), tmp_path / "estimate.xlsx")
    result = recalculate(path)
    assert result.ok, result.errors
    assert result.engine == "python-calibrated"
    assert result.controls_passed >= 5
    assert not result.native_opened
    assert result.cached_values_count == result.formula_count
    assert result.warnings


@pytest.mark.parametrize("formula,error", [("=1/0", "#DIV/0!"), ("=A1", "circular"), ("=UNKNOWN(1)", "unsupported"), ("=Missing!A1", "#REF!")])
def test_fallback_refuses_errors_and_leaves_original_unchanged(tmp_path, monkeypatch, formula, error):
    import tooling.recalculate as module
    monkeypatch.setattr(module, "_select_excel_backend", lambda: None)
    monkeypatch.setattr(module, "_find_libreoffice", lambda: None)
    path = tmp_path / "invalid.xlsx"
    wb = Workbook()
    wb.active["A1"] = formula
    wb.save(path)
    before = path.read_bytes()
    result = recalculate(path)
    assert not result.ok
    assert any(error in item for item in result.errors)
    assert path.read_bytes() == before


def test_fallback_failure_to_calibrate_blocks_delivery(valid_workbook, monkeypatch):
    import tooling.recalculate as module
    monkeypatch.setattr(module, "_select_excel_backend", lambda: None)
    monkeypatch.setattr(module, "_find_libreoffice", lambda: None)
    monkeypatch.setattr(module, "_calibrate_python", lambda: (0, ["control failed"]))
    before = valid_workbook.read_bytes()
    result = recalculate(valid_workbook)
    assert not result.ok
    assert "control failed" in result.errors
    assert valid_workbook.read_bytes() == before


def test_validator_explicit_zero_counts_for_formula_free_document(tmp_path):
    wb = Workbook()
    path = tmp_path / "empty.xlsx"
    wb.save(path)
    report = validate_workbook(path)
    assert report.counts["formulas"] == 0
    assert report.counts["formula_errors"] == 0
    assert report.counts["missing_cached_values"] == 0
    assert report.counts["table_filter_overlaps"] == 0


def test_validator_rejects_stale_cached_total(valid_workbook):
    with ZipFile(valid_workbook) as archive:
        members = [(item, archive.read(item.filename)) for item in archive.infolist()]
    with ZipFile(valid_workbook, "w") as archive:
        for item, data in members:
            if item.filename == "xl/worksheets/sheet12.xml":
                root = ET.fromstring(data)
                for cell in root.iter("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c"):
                    if cell.attrib["r"] == "B6":
                        cell.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}v").text = "999"
                data = ET.tostring(root)
            archive.writestr(item, data)
    report = validate_workbook(valid_workbook)
    assert not report.ok
    assert report.counts["missing_cached_values"] == 0
    assert any("11_Synthese!B6" in error and "reconciliation" in error for error in report.errors)


def test_validator_rejects_changed_formula_with_correct_old_cache_without_mutation(valid_workbook):
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with ZipFile(valid_workbook) as archive:
        members = [(item, archive.read(item.filename)) for item in archive.infolist()]
    with ZipFile(valid_workbook, "w") as archive:
        for item, data in members:
            if item.filename == "xl/worksheets/sheet12.xml":
                root = ET.fromstring(data)
                for cell in root.iter(f"{namespace}c"):
                    if cell.attrib["r"] == "B6":
                        assert cell.find(f"{namespace}v").text == "14"
                        cell.find(f"{namespace}f").text = "1+1"
                data = ET.tostring(root)
            archive.writestr(item, data)
    before = valid_workbook.read_bytes()
    report = validate_workbook(valid_workbook)
    assert not report.ok
    assert report.counts["missing_cached_values"] == 0
    assert report.counts["reconciliation_errors"] == 0
    assert report.counts["formula_semantic_errors"] == 1
    assert any("11_Synthese!B6" in error and "formula/cache" in error for error in report.errors)
    assert valid_workbook.read_bytes() == before


def test_cost_reconciliation_groups_provider_environment_and_scenario(tmp_path):
    data = json.loads(FIXTURE.read_text(encoding="utf-8"))
    data["cloud_costs"].append(dict(data["cloud_costs"][0], id="COST-04", environment="Development", quantity=3, unit_price=2, monthly_multiplier=4))
    data["cloud_costs"].append(dict(data["cloud_costs"][0], id="COST-05", scenario="Economique", quantity=1, unit_price=7))
    brief = tmp_path / "brief.json"
    brief.write_text(json.dumps(data), encoding="utf-8")
    path = build_workbook(load_model(brief), tmp_path / "estimate.xlsx")
    assert recalculate(path).ok
    report = validate_workbook(path)
    assert report.ok, report.errors
    totals = {(item["provider"], item["environment"], item["scenario"]): item["monthly"] for item in report.cost_totals}
    assert totals[("Provider A", "Development", "Standard")] == 24
    assert totals[("Provider A", "Production", "Standard")] == 90
    assert totals[("Provider A", "Production", "Economique")] == 7
    assert load_workbook(path, data_only=True)["11_Synthese"]["B11"].value == 114


def test_native_excel_takes_priority_and_records_verified_opening(valid_workbook, monkeypatch):
    import tooling.recalculate as module
    native_bytes = valid_workbook.read_bytes()
    mutate(valid_workbook, "11_Synthese", "B6", "=SUMIF(Taches[Nature],\"Build\",Taches[JH retenu])")
    monkeypatch.setattr(module, "_select_excel_backend", lambda: "powershell")
    monkeypatch.setattr(module, "_native_excel", lambda candidate, backend: candidate.write_bytes(native_bytes))
    monkeypatch.setattr(module, "_verify_native_final", lambda candidate, backend: None)
    monkeypatch.setattr(module, "_find_libreoffice", lambda: Path("unused-soffice"))
    result = recalculate(valid_workbook)
    assert result.ok
    assert result.engine == "excel-com"
    assert result.native_opened
    assert result.cached_values_count > 40


def test_native_excel_com_roundtrip_when_available(tmp_path):
    import tooling.recalculate as module
    backend = module._select_excel_backend()
    if backend is None:
        pytest.skip("Excel COM is not available on this host")
    path = build_workbook(load_model(FIXTURE), tmp_path / "native.xlsx")
    result = recalculate(path)
    assert result.ok, result.errors
    assert result.engine == "excel-com"
    assert result.native_opened
    assert result.cached_values_count == result.formula_count
    report = validate_workbook(path)
    assert report.ok, report.errors


def test_excel_backend_selects_powershell_when_pywin32_is_missing(monkeypatch):
    import tooling.recalculate as module
    monkeypatch.setattr(module.sys, "platform", "win32")
    monkeypatch.setitem(sys.modules, "win32com", None)
    calls = []
    monkeypatch.setattr(module, "_run_excel_powershell", lambda script, timeout: calls.append(timeout))
    assert module._select_excel_backend() == "powershell"
    assert calls == [30]


def test_failed_native_excel_uses_verified_python_fallback(valid_workbook, monkeypatch):
    import tooling.recalculate as module
    def unavailable(candidate, backend):
        raise OSError("native application unavailable")
    monkeypatch.setattr(module, "_select_excel_backend", lambda: "powershell")
    monkeypatch.setattr(module, "_native_excel", unavailable)
    monkeypatch.setattr(module, "_find_libreoffice", lambda: None)
    result = recalculate(valid_workbook)
    assert result.ok
    assert result.engine == "python-calibrated"
    assert any("native application unavailable" in warning for warning in result.warnings)


def test_libreoffice_success_after_excel_is_unavailable(valid_workbook, monkeypatch):
    import tooling.recalculate as module
    cached_bytes = valid_workbook.read_bytes()
    mutate(valid_workbook, "11_Synthese", "B6", "=1+1")
    events = []
    monkeypatch.setattr(module, "_select_excel_backend", lambda: None)
    monkeypatch.setattr(module, "_find_libreoffice", lambda: Path("soffice"))
    def libreoffice(candidate, executable, directory):
        events.append("libreoffice")
        candidate.write_bytes(cached_bytes)
    monkeypatch.setattr(module, "_native_libreoffice", libreoffice)
    result = recalculate(valid_workbook)
    assert result.ok, result.errors
    assert result.engine == "libreoffice"
    assert result.native_opened
    assert events == ["libreoffice"]


def test_excel_failure_then_libreoffice_success(valid_workbook, monkeypatch):
    import tooling.recalculate as module
    cached_bytes = valid_workbook.read_bytes()
    mutate(valid_workbook, "11_Synthese", "B6", "=1+1")
    events = []
    monkeypatch.setattr(module, "_select_excel_backend", lambda: "powershell")
    monkeypatch.setattr(module, "_find_libreoffice", lambda: Path("soffice"))
    def excel(candidate, backend):
        events.append("excel")
        raise OSError("Excel failed")
    def libreoffice(candidate, executable, directory):
        events.append("libreoffice")
        candidate.write_bytes(cached_bytes)
    monkeypatch.setattr(module, "_native_excel", excel)
    monkeypatch.setattr(module, "_native_libreoffice", libreoffice)
    result = recalculate(valid_workbook)
    assert result.ok, result.errors
    assert result.engine == "libreoffice"
    assert result.native_opened
    assert events == ["excel", "libreoffice"]
    assert any("Excel failed" in warning for warning in result.warnings)


def test_empty_optional_sections_and_zero_effort_still_validate(tmp_path):
    data = json.loads(FIXTURE.read_text(encoding="utf-8"))
    data["tasks"][0].update(jh_low=0, jh_likely=0, jh_high=0)
    for section in ("inventory", "architecture", "cloud_costs", "risks", "decisions", "sources"):
        data[section] = []
    brief = tmp_path / "empty.json"
    brief.write_text(json.dumps(data), encoding="utf-8")
    path = build_workbook(load_model(brief), tmp_path / "empty.xlsx")
    result = recalculate(path)
    assert result.ok, result.errors
    report = validate_workbook(path, diagram_inventory_ids=[])
    assert report.ok, report.errors
    assert report.cost_totals == []


def test_calibration_runs_the_deliberately_broken_controls():
    from tooling.recalculate import _calibrate_python
    passed, errors = _calibrate_python()
    assert errors == []
    assert passed == 10


def test_fallback_numeric_zero_criterion_does_not_match_blanks(tmp_path, monkeypatch):
    import tooling.recalculate as module
    monkeypatch.setattr(module, "_select_excel_backend", lambda: None)
    monkeypatch.setattr(module, "_find_libreoffice", lambda: None)
    wb = Workbook()
    ws = wb.active
    ws.append([0, 5])
    ws.append([1, 7])
    ws.append([None, 11])
    ws["C1"] = '=SUMIF(A1:A3,"0",B1:B3)'
    path = tmp_path / "zero.xlsx"
    wb.save(path)
    result = recalculate(path)
    assert result.ok, result.errors
    assert load_workbook(path, data_only=True).active["C1"].value == 5
