"""Fresh synthetic reproductions of the first independent forward-test defects."""

from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import re
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.worksheet.table import Table
import pytest

from tooling.anonymisation import scan_tree
from tooling.diagram_renderer import render_diagram
from tooling.model import load_model
import tooling.recalculate as calculation
from tooling.workbook_builder import build_workbook
from tooling.workbook_validator import validate_workbook

ROOT = Path(__file__).parents[1]
FIXTURE = ROOT / "tests/fixtures/synthetic-brief.json"
SVG = {"s": "http://www.w3.org/2000/svg"}


@pytest.mark.parametrize("rate", [None, "A chiffrer"])
def test_unknown_rate_preserves_technical_effort_and_blocks_only_finance(tmp_path, monkeypatch, rate):
    monkeypatch.setattr(calculation, "_select_excel_backend", lambda: None)
    monkeypatch.setattr(calculation, "_find_libreoffice", lambda: None)
    data = json.loads(FIXTURE.read_text(encoding="utf-8"))
    assumption = next(row for row in data["assumptions"] if row["key"] == "taux_journalier")
    assumption.update(value=rate, status="A confirmer")
    brief = tmp_path / "unknown.json"
    brief.write_text(json.dumps(data), encoding="utf-8")
    path = build_workbook(load_model(brief), tmp_path / "unknown.xlsx")
    result = calculation.recalculate(path)
    assert result.ok, result.errors
    values = load_workbook(path, data_only=True)
    assert values["11_Synthese"]["B6"].value == 14
    assert values["11_Synthese"]["B7"].value == 1
    assert values["11_Synthese"]["B8"].value is None
    assert values["11_Synthese"]["B9"].value is None
    status = next(row[1].value for row in values["11_Synthese"] if row[0].value == "Statut finance")
    assert status == "Bloque - taux journalier A chiffrer"
    report = validate_workbook(path)
    assert report.ok, report.errors
    assert report.finance_status == "blocked"
    assert report.counts["missing_cached_values"] == 0
    assert report.counts["reconciliation_errors"] == 0
    assert report.warnings
    # An unknown cost is not interchangeable with a zero cache, even if effort is zero.
    wb = load_workbook(path)
    wb["11_Synthese"]["B8"] = "=0"
    wb.save(path)
    assert calculation.recalculate(path).ok
    assert not validate_workbook(path).ok


def _rewrite(path, transform):
    with ZipFile(path) as archive:
        members = [(item, archive.read(item.filename)) for item in archive.infolist()]
    with ZipFile(path, "w") as archive:
        for item, raw in members:
            archive.writestr(item, transform(item.filename, raw))


def test_native_final_publication_strips_private_metadata_without_namespace_or_cache_loss(tmp_path, monkeypatch):
    monkeypatch.setattr(calculation, "_select_excel_backend", lambda: None)
    monkeypatch.setattr(calculation, "_find_libreoffice", lambda: None)
    path = build_workbook(load_model(FIXTURE), tmp_path / "published.xlsx")
    assert calculation.recalculate(path).ok
    cached = path.read_bytes()
    private_path = "C:" + chr(92) + chr(92).join(["Users", "private-user", "work"])

    def native(candidate, backend):
        candidate.write_bytes(cached)
        def metadata(name, raw):
            if name == "xl/workbook.xml":
                raw = raw.replace(b"</workbook>", ('<extLst><ext uri="privacy-test"><x15ac:absPath xmlns:x15ac="http://schemas.microsoft.com/office/spreadsheetml/2010/11/ac" url="' + private_path + '"/></ext></extLst></workbook>').encode())
                if b"absPath" not in raw:
                    # The fallback serializer uses a namespace prefix.
                    root = ET.fromstring(raw)
                    ET.SubElement(root, "{http://schemas.microsoft.com/office/spreadsheetml/2010/11/ac}absPath", url=private_path)
                    raw = ET.tostring(root)
            if name == "docProps/core.xml":
                raw = re.sub(rb"(<(?:\w+:)?creator[^>]*>).*?(</(?:\w+:)?creator>)", rb"\1private-author\2", raw)
            return raw
        _rewrite(candidate, metadata)

    monkeypatch.setattr(calculation, "_select_excel_backend", lambda: "powershell")
    monkeypatch.setattr(calculation, "_native_excel", native)
    monkeypatch.setattr(calculation, "_verify_native_final", lambda candidate, backend: None, raising=False)
    result = calculation.recalculate(path)
    assert result.ok, result.errors
    with ZipFile(path) as archive:
        for name in archive.namelist():
            if name.endswith((".xml", ".rels")):
                raw = archive.read(name)
                ET.fromstring(raw)
                assert b"absPath" not in raw
                assert b"private-author" not in raw
                assert private_path.encode() not in raw
        core = archive.read("docProps/core.xml")
        assert b"dcterms:W3CDTF" in core
        assert b"xmlns:dcterms=" in core
    assert validate_workbook(path).ok
    assert result.publication_inspected_parts > 12
    assert load_workbook(path, data_only=True)["11_Synthese"]["B6"].value == 14


def test_office_structural_guid_is_not_a_cloud_account_but_cell_guid_still_is(tmp_path):
    suffix = "0" * 12
    guid = "{00000000-0000-0000-0000-" + suffix + "}"
    path = tmp_path / "structural.xlsx"
    with ZipFile(path, "w") as archive:
        archive.writestr("xl/worksheets/sheet1.xml", '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:xr="http://schemas.microsoft.com/office/spreadsheetml/2014/revision" xr:uid="' + guid + '"><sheetData/></worksheet>')
    assert scan_tree(tmp_path).clean
    with ZipFile(path, "a") as archive:
        archive.writestr("xl/sharedStrings.xml", '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><t>' + guid + '</t></si></sst>')
    assert not scan_tree(tmp_path).clean


def test_excel_revision_document_id_numeric_guid_is_structural(tmp_path):
    # Catches treating revisionPtr.documentId as a cloud account, or globally
    # exempting its GUID when the same identifier occurs in business cells.
    guid = "{00000000-0000-0000-0000-" + "0" * 12 + "}"
    path = tmp_path / "revision.xlsx"
    with ZipFile(path, "w") as archive:
        archive.writestr("xl/workbook.xml", '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:xr="http://schemas.microsoft.com/office/spreadsheetml/2014/revision"><xr:revisionPtr documentId="13_ncr:1_' + guid + '"/></workbook>')
    original = path.read_bytes()
    report = scan_tree(tmp_path)
    assert report.clean, report.summary()
    assert path.read_bytes() == original
    with ZipFile(path, "a") as archive:
        archive.writestr("xl/sharedStrings.xml", '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><t>' + guid + '</t></si></sst>')
    report = scan_tree(tmp_path)
    assert not report.clean
    assert {(item.path, item.rule) for item in report.matches} == {
        ("revision.xlsx!xl/sharedStrings.xml", "cloud-account-id"),
    }


@pytest.mark.parametrize("attribute,clean", [
    ('documentId="13_ncr:1_GUID"', True),
    ('private-documentId="13_ncr:1_GUID"', False),
    ('note="documentId=\'13_ncr:1_GUID\'"', False),
    ('note="business documentId=\'13_ncr:1_GUID\' > additional data"', False),
    ('p:documentId="13_ncr:1_GUID"', False),
], ids=["exact", "hyphenated-name", "inside-note", "quoted-tag-delimiter", "prefixed"])
def test_revision_document_id_exemption_requires_actual_unqualified_attribute(tmp_path, attribute, clean):
    # Catches matching documentId inside another attribute's name or value.
    guid = "{00000000-0000-0000-0000-" + "0" * 12 + "}"
    path = tmp_path / "revision.xlsx"
    with ZipFile(path, "w") as archive:
        archive.writestr("xl/workbook.xml", '<xr:revisionPtr xmlns:xr="http://schemas.microsoft.com/office/spreadsheetml/2014/revision" xmlns:p="urn:synthetic" ' + attribute.replace("GUID", guid) + '/>')
    original = path.read_bytes()
    report = scan_tree(tmp_path)
    assert report.clean is clean, report.summary()
    assert path.read_bytes() == original
    if not clean:
        assert {(item.path, item.rule) for item in report.matches} == {
            ("revision.xlsx!xl/workbook.xml", "cloud-account-id"),
        }


@pytest.mark.parametrize("quote", ['"', "'"], ids=["double-quoted", "single-quoted"])
@pytest.mark.parametrize("note_first", [False, True], ids=["document-first", "note-first"])
def test_document_id_exemption_does_not_mask_note_with_quoted_tag_delimiter(tmp_path, quote, note_first):
    # Catches a truncated start tag granting the real ID's exemption to note text.
    guid = "{00000000-0000-0000-0000-" + "0" * 12 + "}"
    value = "13_ncr:1_" + guid
    inner_quote = "'" if quote == '"' else '"'
    structural_guid = "{00000000-0000-0000-0000-" + "a" * 12 + "}"
    document_id = "documentId=" + quote + "13_ncr:1_" + structural_guid + quote
    note = "note=" + quote + "business documentId=" + inner_quote + value + inner_quote + " > additional data" + quote
    attributes = [note, document_id] if note_first else [document_id, note]
    path = tmp_path / "revision.xlsx"
    with ZipFile(path, "w") as archive:
        archive.writestr("xl/workbook.xml", '<xr:revisionPtr xmlns:xr="http://schemas.microsoft.com/office/spreadsheetml/2014/revision" ' + " ".join(attributes) + '/>')
    original = path.read_bytes()
    report = scan_tree(tmp_path)
    assert not report.clean, report.summary()
    assert {(item.path, item.rule) for item in report.matches} == {
        ("revision.xlsx!xl/workbook.xml", "cloud-account-id"),
    }
    assert path.read_bytes() == original


@pytest.mark.parametrize("quote", ['"', "'"], ids=["double-quoted", "single-quoted"])
@pytest.mark.parametrize("part,start,attribute", [
    ("xl/worksheets/sheet1.xml", '<worksheet xmlns:xr="http://schemas.microsoft.com/office/spreadsheetml/2014/revision"', "xr:uid"),
    ("xl/workbook.xml", '<xr:revisionPtr xmlns:xr="http://schemas.microsoft.com/office/spreadsheetml/2014/revision" xmlns:xr10="http://schemas.microsoft.com/office/spreadsheetml/2016/revision10"', "xr10:uidLastSave"),
    ("xl/drawings/drawing1.xml", '<a16:creationId xmlns:a16="http://schemas.microsoft.com/office/drawing/2014/main"', "id"),
], ids=["revision-uid", "revision-save", "drawing-id"])
def test_structural_guid_exemption_does_not_mask_embedded_note_attribute(tmp_path, quote, part, start, attribute):
    # Catches applying a structural GUID exemption to a substring of note text.
    guid = "{00000000-0000-0000-0000-" + "0" * 12 + "}"
    inner_quote = "'" if quote == '"' else '"'
    xml = start + " note=" + quote + "business > " + attribute + "=" + inner_quote + guid + inner_quote + quote + "/>"
    path = tmp_path / "embedded.xlsx"
    with ZipFile(path, "w") as archive:
        archive.writestr(part, xml)
    original = path.read_bytes()
    report = scan_tree(tmp_path)
    assert not report.clean, report.summary()
    assert {(item.path, item.rule) for item in report.matches} == {
        ("embedded.xlsx!" + part, "cloud-account-id"),
    }
    assert path.read_bytes() == original


@pytest.mark.parametrize("part,xml", [
    ("xl/worksheets/sheet1.xml", '<worksheet xmlns:r="http://schemas.microsoft.com/office/spreadsheetml/2015/revision2" r:uid="GUID"/>'),
    ("xl/worksheets/sheet1.xml", '<worksheet xmlns:r="http://schemas.microsoft.com/office/spreadsheetml/2016/revision3" r:uid="GUID"/>'),
    ("xl/workbook.xml", '<xr:revisionPtr xmlns:xr="http://schemas.microsoft.com/office/spreadsheetml/2014/revision" xmlns:xr10="http://schemas.microsoft.com/office/spreadsheetml/2016/revision10" xr10:uidLastSave="GUID"/>'),
    ("xl/drawings/drawing1.xml", '<a16:creationId xmlns:a16="http://schemas.microsoft.com/office/drawing/2014/main" id="GUID"/>'),
])
def test_native_office_revision_and_creation_guid_context(tmp_path, part, xml):
    guid = "{00000000-0000-0000-0000-" + "0" * 12 + "}"
    with ZipFile(tmp_path / "native.xlsx", "w") as archive:
        archive.writestr(part, xml.replace("GUID", guid))
    assert scan_tree(tmp_path).clean


def _segments(path):
    numbers = list(map(float, re.findall(r"-?\d+(?:\.\d+)?", path.attrib["d"])))
    points = list(zip(numbers[::2], numbers[1::2]))
    return list(zip(points, points[1:]))


def _shared(first, second):
    (ax, ay), (bx, by) = first
    (cx, cy), (dx, dy) = second
    return ((ay == by == cy == dy and max(min(ax, bx), min(cx, dx)) < min(max(ax, bx), max(cx, dx)))
            or (ax == bx == cx == dx and max(min(ay, by), min(cy, dy)) < min(max(ay, by), max(cy, dy))))


def test_forward_three_zone_chain_has_separate_flow_lanes_and_label_boxes(tmp_path):
    graph = {"nodes": [{"id": key, "component": key, "environment": "Production", "zone": zone}
                       for key, zone in [("EDGE", "Edge"), ("APP", "Application"), ("DATA", "Data")]],
             "edges": [{"id": "F-01", "source": "EDGE", "target": "APP", "label": "HTTPS"},
                       {"id": "F-02", "source": "APP", "target": "DATA", "label": "Connexion base chiffree"}]}
    render_diagram(graph, tmp_path / "a.svg", tmp_path / "a.png")
    svg = ET.parse(tmp_path / "a.svg")
    paths = svg.findall(".//s:path[@data-edge-id]", SVG)
    assert not any(_shared(a, b) for a in _segments(paths[0]) for b in _segments(paths[1]))
    labels = svg.findall(".//s:rect[@data-edge-label-id]", SVG)
    assert len(labels) == 2
    boxes = [(float(r.attrib["x"]), float(r.attrib["y"]), float(r.attrib["width"]), float(r.attrib["height"])) for r in labels]
    (x, y, w, h), (a, b, c, d) = boxes
    assert max(x, a) >= min(x + w, a + c) or max(y, b) >= min(y + h, b + d)
    from tooling.diagram_renderer import _clear_segment
    node_boxes = []
    for node in svg.findall(".//s:g[@data-node-id]/s:rect", SVG):
        nx, ny, nw, nh = [float(node.attrib[key]) for key in ("x", "y", "width", "height")]
        node_boxes.append((nx, ny, nx + nw, ny + nh))
    for x, y, w, h in boxes:
        rect = (x, y, x + w, y + h)
        assert not any(max(x, a) < min(x + w, c) and max(y, b) < min(y + h, d) for a, b, c, d in node_boxes)
        assert all(_clear_segment(a, b, [rect]) for path in paths for a, b in _segments(path))
    assert {t.text for t in svg.findall(".//s:text", SVG)} >= {"HTTPS", "Connexion base chiffree"}
    shuffled = deepcopy(graph)
    shuffled["nodes"].reverse()
    shuffled["edges"].reverse()
    render_diagram(shuffled, tmp_path / "b.svg", tmp_path / "b.png")
    assert (tmp_path / "a.svg").read_bytes() == (tmp_path / "b.svg").read_bytes()
    assert (tmp_path / "a.png").read_bytes() == (tmp_path / "b.png").read_bytes()


def test_scenario_annex_headers_do_not_expand_canonical_a_to_f(tmp_path, monkeypatch):
    monkeypatch.setattr(calculation, "_select_excel_backend", lambda: None)
    monkeypatch.setattr(calculation, "_find_libreoffice", lambda: None)
    path = build_workbook(load_model(FIXTURE), tmp_path / "annex.xlsx")
    wb = load_workbook(path)
    ws = wb["06_Scenarios"]
    ws["H4"], ws["I4"], ws["H5"], ws["I5"] = "Dimension", "Capacite", "CPU", 4
    ws.add_table(Table(displayName="CapacityAnnex", ref="H4:I5"))
    wb.save(path)
    assert calculation.recalculate(path).ok
    report = validate_workbook(path)
    assert report.ok, report.errors
    assert report.counts["annex_blocks"] == 1
    # Separate annex validation must reject a missing header, not ignore it.
    wb = load_workbook(path)
    wb["06_Scenarios"]["I4"] = None
    wb.save(path)
    assert calculation.recalculate(path).ok
    report = validate_workbook(path)
    assert not report.ok
    assert any("annex" in e.lower() for e in report.errors)


@pytest.mark.parametrize("rate", [None, "A chiffrer"])
def test_unknown_rate_native_final_bytes_and_capacity_formula(tmp_path, rate):
    if calculation._select_excel_backend() is None:
        pytest.skip("Excel COM unavailable")
    data = json.loads(FIXTURE.read_text(encoding="utf-8"))
    next(row for row in data["assumptions"] if row["key"] == "taux_journalier")["value"] = rate
    brief = tmp_path / "unknown.json"
    brief.write_text(json.dumps(data), encoding="utf-8")
    path = build_workbook(load_model(brief), tmp_path / "native.xlsx")
    wb = load_workbook(path)
    ws = wb["06_Scenarios"]
    for cell, value in {"H4": "Demand", "I4": "Unit capacity", "J4": "Units", "H5": 6, "I5": 3, "J5": "=H5/I5"}.items():
        ws[cell] = value
    ws.add_table(Table(displayName="SyntheticCapacity", ref="H4:J5"))
    wb.save(path)
    result = calculation.recalculate(path)
    assert result.ok and result.engine == "excel-com", result.errors
    assert result.final_native_opened and result.publication_inspected_parts > 12
    report = validate_workbook(path, asserted_totals={"06_Scenarios!J5": 2, "11_Synthese!B6": 14})
    assert report.ok and report.finance_status == "blocked", report.errors
    values = load_workbook(path, data_only=True)
    assert values["11_Synthese"]["B8"].value is None
    assert values["11_Synthese"]["B9"].value is None
    assert scan_tree(tmp_path).clean


def test_publication_refuses_private_relationship_without_altering_original(tmp_path, monkeypatch):
    monkeypatch.setattr(calculation, "_select_excel_backend", lambda: None)
    monkeypatch.setattr(calculation, "_find_libreoffice", lambda: None)
    path = build_workbook(load_model(FIXTURE), tmp_path / "private.xlsx")
    private = "file:///" + "C:/" + "/".join(["Users", "private-user", "private.xlsx"])
    def relationship(name, raw):
        if name == "xl/_rels/workbook.xml.rels":
            root = ET.fromstring(raw)
            ET.SubElement(root, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                          Id="PrivatePath", Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", Target=private, TargetMode="External")
            return ET.tostring(root)
        return raw
    _rewrite(path, relationship)
    original = path.read_bytes()
    result = calculation.recalculate(path)
    assert not result.ok and any("local-user path" in error for error in result.errors)
    assert path.read_bytes() == original
