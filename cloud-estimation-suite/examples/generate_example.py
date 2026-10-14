"""Generate a public synthetic model, diagram and recalculated XLSX acceptance case.

This example reads only the repository's public synthetic fixture. Prices and
capacity are invented assumptions; no cloud calls or private inputs are used.
"""

import argparse
from copy import deepcopy
import json
from pathlib import Path
import re
import sys
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from tooling.anonymisation import scan_tree
from tooling.diagram_renderer import render_diagram
from tooling.model import load_model
from tooling.recalculate import recalculate
from tooling.publication import sanitize_publication
from tooling.workbook_builder import build_workbook
from tooling.workbook_validator import validate_workbook


def synthetic_model():
    data = json.loads((ROOT / "tests/fixtures/synthetic-brief.json").read_text(encoding="utf-8"))
    data["metadata"]["notes"] = (
        "Entirely synthetic public acceptance example, not a quote. Standard selected; "
        "Provider A is a provisional illustration. Build JH and Run JH/month are separate. "
        "Provider alternatives must not be added. Forecast FinOps savings require measurement. "
        "Kubernetes: Non applicable after suitability review; managed runtime selected. "
        "Taxes, discounts, FX and growth excluded; non-cloud OPEX assumed zero. "
        "Capacity overloads and delivery overload remain open; this is not an approved production design."
    )
    data["tasks"][0]["coverage"] = "Partial, evidenced synthetic coverage"
    data["tasks"][0]["notes"] = "Gross 3/4/6 JH; covered 1/1/1 JH (SYN-02); net 2/3/5 JH. One canonical owner; no aliases."
    data["tasks"][3]["acceptance_criterion"] = "Staging harness records base/high CPU, memory, storage, IOPS, bandwidth and concurrency; verify target and overload before approval"
    data["tasks"][3]["expected_evidence"] = "Synthetic load assumptions now; measured traces, latency percentiles and saturation report required for approval"
    data["tasks"][4]["notes"] = "Synthetic transition only. Preparation, transfer, integrity, cutover, rollback and owner-only decommissioning in runbook; hypercare is a separate recurring operation. No live source access."
    additions = [
        ("TASK-07", "Operations", "finops", "Schedule non-production runtime", "Reversible schedule", ["TASK-05"], (1, 2, 3), "Forecast; compare metered usage before/after at equal scope; restore always-on schedule on validation failure"),
        ("TASK-08", "Delivery", "architecture-diagram", "Reconcile inventory and explicit flows", "Embedded diagram and flow register", ["TASK-03"], (0.5, 1, 2), "Each node matches inventory; HTTPS edge matches SYN-03; no inferred connection"),
        ("TASK-09", "Foundation", "cloud-architecture", "Resolve managed runtime versus cluster suitability", "Architecture decision", ["TASK-01"], (0.5, 1, 2), "Confirm application isolation, backup and recovery evidence; cluster work excluded because no platform requirement"),
        ("TASK-10", "Foundation", "provider-selection", "Compare equal service scope and three-year TCO", "Provider comparison", ["TASK-01"], (0.5, 1, 2), "Nine equal-scope alternatives, regional eligibility remains hypothetical; do not approve without official dated quotes"),
    ]
    for task_id, lot, domain, task, deliverable, dependencies, estimates, criterion in additions:
        record = deepcopy(data["tasks"][0])
        record.update(id=task_id, lot=lot, domain=domain, task=task, deliverable=deliverable,
                      dependencies=dependencies, acceptance_criterion=criterion,
                      expected_evidence=deliverable + " with SYN-01 assumptions", scenario_override=None,
                      jh_low=estimates[0], jh_likely=estimates[1], jh_high=estimates[2], coverage="Full",
                      notes="Synthetic net additional Build; one owner and no duplicate deliverable.")
        data["tasks"].append(record)
    data["inventory"].append(dict(id="COMP-04", component="Managed data and platform services", environment="Production",
                                  quantity=1, unit="platform envelope", notes="Usable data, backups, network/egress, logging, registry/build and support; synthetic envelope pricing, not a SKU quote."))
    data["architecture"][0]["notes"] = "Provisional Standard: two runtime replicas. Economique: one, Resilient: three. High workload overload remains open. See DEC-CONFLICT and capacity assumptions."
    data["architecture"].append(dict(id="ARCH-02", component="Managed data and platform services", type="Managed backend",
                                     environment="Production", provider="Provider A", decision="Managed backend, separate shared capacity",
                                     notes="COMP-04; shared data capacity is not multiplied by runtime replicas. Backup and restore evidence required."))
    data["cloud_costs"] = []
    for provider, price in (("Provider A", 45), ("Provider B", 50), ("Provider C", 55)):
        for scenario, replicas, envelope in (("Economique", 1, 24), ("Standard", 2, 36), ("Resilient", 3, 60)):
            for service, environment, quantity, unit_price, multiplier in (
                ("Managed runtime", "Production", replicas, price, 1),
                ("Managed runtime scheduled (Development + Staging)", "Non-production", 2, price, 0.5),
                ("Managed data, backups, network/egress, logging, build/registry, support", "Production", 1, envelope, 1),
            ):
                data["cloud_costs"].append(dict(id=f"COST-{len(data['cloud_costs']) + 1:02}", provider=provider, service=service,
                                               environment=environment, scenario=scenario, unit="unit/month", quantity=quantity,
                                               unit_price=unit_price, currency="EUR", monthly_multiplier=multiplier,
                                               source="SYN-01", confidence="Low", notes="Synthetic forecast only; same service perimeter, alternative provider. Replace envelope by verified SKU breakdown before real use."))

    def assumption(key, value, unit, notes):
        data["assumptions"].append(dict(key=key, value=value, unit=unit, status="A confirmer", source="SYN-01", notes=notes))

    for key, value, unit, notes in [
        ("hours_per_jh", 8, "h/JH", "Editable convention; task estimates already in JH."),
        ("schedule.foundation_available", 6, "JH", "Period 1; collective capacity across Foundation roles, not per-role capacity. Owner: delivery lead. Dependency TASK-01 before TASK-02/09/10; individual role availability remains unverified."),
        ("schedule.foundation_overload", 0, "JH", "Dynamic formula: Foundation Build minus available JH; positive = overload, DEC-SCHEDULE."),
        ("capacity.production.Standard.base.cpu", 6, "vCPU", "Assumed aligned CPU demand; independent service scenario and workload band."),
        ("capacity.production.Standard.high.cpu", 12, "vCPU", "Assumed high-band CPU demand, not a benchmark."),
        ("capacity.production.Standard.base.memory", 8, "GiB", "Synthetic working-set demand; per runtime unit 8 GiB."),
        ("capacity.production.Standard.base.storage_gib", 100, "GiB", "Shared usable data demand; backend capacity 200 GiB, not replica-scaled."),
        ("capacity.production.Standard.base.iops", 1000, "operations/s", "Shared backend assumed effective 2000 operations/s; latency/mix unmeasured."),
        ("capacity.production.Standard.base.bandwidth", 100, "Mbit/s", "Shared link assumed 200 Mbit/s effective; egress envelope is separate pricing."),
        ("capacity.production.Standard.base.concurrency", 100, "requests", "Runtime assumed safe 100 concurrent requests/unit; no throughput inferred."),
        ("capacity.production.Standard.base.utilisation_target", 0.75, "ratio", "CPU/memory/concurrency target, includes headroom; shared backend requires independent checks."),
        ("capacity.production.Standard.base.scaling_limit", 3, "instances", "Hypothetical ceiling; does not make high-band demand feasible."),
        ("capacity.production.Standard.base.failure_domain_reserve", 1, "instances", "One runtime loss; does not prove zone/region DR."),
        ("capacity.production.Standard.cpu_per_unit", 4, "vCPU/instance", "Effective capacity assumption, not hardware specification."),
        ("capacity.production.Standard.units", 2, "instances", "Standard runtime count corresponds to COMP-03 and Provider A Standard cost quantity."),
        ("capacity.production.Standard.high.overload_ratio", 0, "ratio", "Dynamic: high demand / effective CPU at target; >1 infeasible. Other dimensions and failure event require measurement."),
        ("contradiction.old_units", 1, "instances", "SYN-02 older single-replica statement; same workload scope."),
        ("contradiction.selected_units", 2, "instances", "SYN-03 later resilience requirement; provisional selected branch in DEC-CONFLICT."),
        ("contradiction.unit_price", 45, "EUR/instance/month", "Provider A synthetic unit price; not an official tariff."),
        ("contradiction.cloud_delta", 0, "EUR/mois", "Dynamic selected minus old runtime cost; excludes common envelope/non-production."),
        ("contradiction.build_old", 1, "JH", "Single-replica infrastructure branch; mutually exclusive, not included as another task."),
        ("contradiction.build_selected", 3, "JH", "Selected infrastructure branch equals TASK-01 Standard net JH."),
        ("contradiction.build_delta", 0, "JH", "Dynamic selected minus old Build."),
        ("contradiction.run_old", 0.5, "JH/mois", "Old branch monthly operations forecast."),
        ("contradiction.run_selected", 1, "JH/mois", "Selected monthly operations equals TASK-06 Standard."),
        ("contradiction.run_delta", 0, "JH/mois", "Dynamic selected minus old monthly Run."),
        ("finops.baseline", 90, "EUR/mois", "Forecast: Development and Staging always-on, two units at Provider A price. No measured billing."),
        ("finops.savings_low_fraction", 0.25, "ratio", "Conservative forecast; retains on-demand wakeup."),
        ("finops.savings_high_fraction", 0.5, "ratio", "Selected illustrative schedule; recurring costs already include this schedule, do not subtract savings twice."),
        ("finops.savings_low", 0, "EUR/mois", "Dynamic baseline times low saving fraction."),
        ("finops.savings_high", 0, "EUR/mois", "Dynamic baseline times high saving fraction."),
        ("finops.savings_selected", 0, "EUR/mois", "Dynamic baseline minus scheduled Provider A Standard non-production charges."),
        ("tco_horizon_years", 3, "years", "Equal scope and constant demand; taxes, FX, discounts and growth excluded."),
        ("non_cloud_opex", 0, "EUR/mois", "Explicit synthetic absence, not an unknown represented as zero."),
        ("comparison_scenario", "Standard", "scenario", "Same reference scenario for all provider TCOs; independent of scenario_global, task overrides still apply."),
        ("tco.ProviderA.Standard", 0, "EUR/3years", "Dynamic selected Build plus 36 months of human Run, Provider A Standard cloud and non-cloud OPEX; alternatives never added."),
        ("tco.ProviderB.Standard", 0, "EUR/3years", "Dynamic same-perimeter Provider B comparison; Build, human Run, cloud and horizon remain separate in Scenarios."),
        ("tco.ProviderC.Standard", 0, "EUR/3years", "Dynamic same-perimeter Provider C comparison; alternatives never added."),
        ("region_eligibility", "A/B/C: hypothetical single-region eligibility; multi-region DR unverified", "scope", "No real provider or regional suitability claim."),
        ("commercial_policy", "No commitments, same currency, invented prices, no discounts", "scope", "Runtime quotes and region eligibility required before approval."),
        ("weighted_criteria", "Operational fit 50%, recovery evidence 30%, TCO 20%; unscored until evidence", "scope", "No invented suitability scores."),
        ("source_write_policy", "owner-only", "policy", "Synthetic migration; generator performs no source/environment mutations."),
    ]:
        assumption(key, value, unit, notes)
    data["decisions"] += [
        dict(id="DEC-CONFLICT", question="SYN-02 single replica conflicts with SYN-03 two replicas", option="Provisional two-replica branch",
             rationale="Owner: architecture lead; resolve before design approval. See contradiction.* calculation cells for Build, monthly Run and cloud deltas. Single-replica branch accepts outage; only selected branch enters tasks. All three service scenarios keep the same perimeter.", status="Provisional", source="SYN-02; SYN-03"),
        dict(id="DEC-FINOPS", question="Reduce forecast non-production idle costs", option="Reversible 50% runtime scheduling forecast",
             rationale="Savings range and selected saving are dynamic finops.* cells. TASK-07 Build and TASK-06 monthly Run. Roll back schedule if wakeup affects testing; meter usage before/after at equal scope. Savings unvalidated; scheduled cost already reflects them.", status="Forecast, unvalidated", source="SYN-01"),
        dict(id="DEC-ARCH", question="Managed runtime or Kubernetes platform", option="Managed runtime and shared managed backend",
             rationale="No cluster requirement in this synthetic scope; Kubernetes work Non applicable. Retain application isolation and restore evidence; managed runtime reduces platform ownership, operator monthly Run remains explicit.", status="Provisional", source="SYN-01"),
        dict(id="DEC-SCHEDULE", question="Foundation Build exceeds Period 1 capacity", option="Move dependent Foundation work into another period before approval",
             rationale="Dynamic schedule.foundation_overload compares all Foundation Build with collective capacity, not a per-role budget. Delivery lead must confirm collective and individual role availability; subsequent periods cannot start before prerequisites. No approved calendar until overload is resolved.", status="Open", source="SYN-01"),
    ]
    data["risks"] += [
        dict(id="RISK-OVERLOAD", risk="High workload CPU exceeds Standard target capacity", impact="Service target infeasible at high band; limited scaling and failure reserve", probability="Assumed", mitigation="TASK-04 measurement; revisit replicas, quotas and placement before approval", owner="Performance engineer", status="Open"),
        dict(id="RISK-SCHEDULE", risk="Foundation delivery overload", impact="Dependent delivery cannot meet provisional period", probability="Certain under forecast", mitigation="DEC-SCHEDULE; role owner replans capacity before approval", owner="Platform engineer", status="Open"),
        dict(id="RISK-FINOPS", risk="Scheduled runtime unavailable during unplanned testing", impact="Wakeup delay", probability="Medium", mitigation="Retain on-demand wakeup and reversible always-on schedule", owner="Operations engineer", status="Open"),
    ]
    data["sources"] += [
        dict(id="SYN-02", title="Earlier synthetic single-replica brief and coverage evidence", type="Synthetic", date="2026-10-01", url="", notes="Invented conflicting requirement and coverage; no private source."),
        dict(id="SYN-03", title="Later synthetic resilience requirement and flow register", type="Synthetic", date="2026-10-07", url="", notes="Invented two-replica requirement; FLOW-01 COMP-03 to COMP-04 HTTPS. No live topology."),
    ]
    return data


def _public_metadata(path):
    """Idempotent shared normalization; structural OOXML GUIDs are retained."""
    return sanitize_publication(path)


def generate(output, require_native=False):
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    data = synthetic_model()
    model_path = output / "synthetic-cloud-brief.json"
    model_path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    model = load_model(model_path)
    architecture = dict(nodes=[dict(row, zone="Application" if row["id"] != "COMP-04" else "Data", provider="Provider A", icon="neutral") for row in model.inventory],
                        edges=[dict(id="FLOW-01", source="COMP-03", target="COMP-04", label="HTTPS (SYN-03)", evidence="SYN-03")])
    (output / "synthetic-architecture.json").write_text(json.dumps(architecture, indent=2) + "\n", encoding="utf-8")
    manifest = render_diagram(architecture, output / "synthetic-architecture.svg", output / "synthetic-architecture.png")
    path = build_workbook(model, output / "synthetic-cloud-estimate.xlsx", architecture=architecture)
    workbook = load_workbook(path)
    assumptions = workbook["01_Hypotheses"]
    cells = {row[0].value: row[1].coordinate for row in assumptions if row[0].value}
    ref = lambda key: cells[key]
    comparison = workbook["06_Scenarios"]
    comparison_header = comparison.max_row + 4
    comparison.cell(comparison_header - 1, 1, "Comparaison fournisseurs - scenario de reference et horizon dans Hypotheses")
    headers = ["Fournisseur", "Build (JH)", "Run humain (JH/mois)", "Cloud (EUR/mois)", "Horizon (annees)", "TCO (EUR)"]
    for column, label in enumerate(headers, 1):
        cell = comparison.cell(comparison_header, column, label)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(wrap_text=True)
    scenario_reference = f"'01_Hypotheses'!{ref('comparison_scenario')}"

    def reference_effort(nature):
        def subtotal(column, override):
            return f'SUMIFS(Taches[{column}],Taches[Nature],"{nature}",Taches[Scenario override],"{override}")'
        baseline = f'IF({scenario_reference}="Economique",{subtotal("JH bas", "")},IF({scenario_reference}="Standard",{subtotal("JH probable", "")},{subtotal("JH haut", "")}))'
        overrides = [subtotal(column, override) for column, override in
                     (("JH bas", "Economique"), ("JH probable", "Standard"), ("JH haut", "Resilient"))]
        return "=" + "+".join([baseline] + overrides)

    comparison_rows = {}
    for row, provider in enumerate(("Provider A", "Provider B", "Provider C"), comparison_header + 1):
        comparison_rows[provider] = row
        comparison.cell(row, 1, provider)
        for column, formula in enumerate([
            reference_effort("Build"), reference_effort("Run"),
            f'=SUMIFS(CoutsCloud[Cout mensuel (EUR)],CoutsCloud[Fournisseur],A{row},CoutsCloud[Scenario],{scenario_reference})',
            f"='01_Hypotheses'!{ref('tco_horizon_years')}",
            f'=IF(taux_journalier="","",IF(taux_journalier="A chiffrer","",B{row}*taux_journalier+(C{row}*taux_journalier+D{row}+\'01_Hypotheses\'!{ref("non_cloud_opex")})*mois_par_an*E{row}))',
        ], 2):
            cell = comparison.cell(row, column, formula)
            cell.font = Font(name="Arial", color="008000")
            cell.number_format = '#,##0.00;(#,##0.00);"-"'
    table = Table(displayName="ComparaisonFournisseurs", ref=f"A{comparison_header}:F{comparison_header + 3}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    comparison.add_table(table)
    comparison.row_dimensions[comparison_header].height = 34
    formulas = {
        "schedule.foundation_overload": f'=SUMIFS(Taches[JH retenu],Taches[Lot],"Foundation",Taches[Nature],"Build")-{ref("schedule.foundation_available")}',
        "capacity.production.Standard.high.overload_ratio": f'={ref("capacity.production.Standard.high.cpu")}/({ref("capacity.production.Standard.cpu_per_unit")}*{ref("capacity.production.Standard.units")}*{ref("capacity.production.Standard.base.utilisation_target")})',
        "contradiction.cloud_delta": f'=({ref("contradiction.selected_units")}-{ref("contradiction.old_units")})*{ref("contradiction.unit_price")}',
        "contradiction.build_delta": f'={ref("contradiction.build_selected")}-{ref("contradiction.build_old")}',
        "contradiction.run_delta": f'={ref("contradiction.run_selected")}-{ref("contradiction.run_old")}',
        "finops.savings_low": f'={ref("finops.baseline")}*{ref("finops.savings_low_fraction")}',
        "finops.savings_high": f'={ref("finops.baseline")}*{ref("finops.savings_high_fraction")}',
        "finops.savings_selected": f'={ref("finops.baseline")}-SUMIFS(CoutsCloud[Cout mensuel (EUR)],CoutsCloud[Fournisseur],"Provider A",CoutsCloud[Scenario],"Standard",CoutsCloud[Environnement],"Non-production")',
    }
    formulas.update({f"tco.Provider{provider[-1]}.Standard": f"='06_Scenarios'!F{row}"
                     for provider, row in comparison_rows.items()})
    for key, formula in formulas.items():
        assumptions[ref(key)] = formula
        assumptions[ref(key)].font = Font(name="Arial", color="008000" if "!" in formula or "Taches[" in formula or "CoutsCloud[" in formula else "000000")
        assumptions[ref(key)].number_format = '#,##0.00;(#,##0.00);"-"'
    workbook.properties.creator = "Public synthetic example"
    workbook.properties.lastModifiedBy = "Public synthetic example"
    workbook.save(path)
    recalculation = recalculate(path)
    if not recalculation.ok or require_native and not recalculation.native_opened:
        raise RuntimeError(json.dumps(recalculation.to_dict()))
    _public_metadata(path)
    # Hand-derived expectations, independent of the workbook calculation engine.
    expectations = dict(zip(formulas, (3, 2, 45, 2, 0.5, 22.5, 45, 45, 41906, 42446, 42986)))
    assertions = {f"01_Hypotheses!{ref(key)}": value for key, value in expectations.items()}
    assertions.update({"11_Synthese!B6": 19, "11_Synthese!B7": 1, "11_Synthese!B8": 12350,
                       "11_Synthese!B9": 650, "11_Synthese!B11": 171,
                       "06_Scenarios!D5": 9.5, "06_Scenarios!D6": 17, "06_Scenarios!D7": 30})
    # Hand-derived: 19*650 + 36*(1*650 + cloud), cloud=171/186/201.
    comparison_expected = {"Provider A": (19, 1, 171, 3, 41906),
                           "Provider B": (19, 1, 186, 3, 42446),
                           "Provider C": (19, 1, 201, 3, 42986)}
    for provider, row in comparison_rows.items():
        assertions.update({f"06_Scenarios!{comparison.cell(row, column).coordinate}": value
                           for column, value in enumerate(comparison_expected[provider], 2)})
    (output / "synthetic-assertions.json").write_text(json.dumps(assertions, indent=2) + "\n", encoding="utf-8")
    report = validate_workbook(path, asserted_totals=assertions, diagram_inventory_ids=manifest.node_ids)
    scan = scan_tree(output)
    if not report.ok or not scan.clean:
        raise RuntimeError(json.dumps(dict(validation=report.to_dict(), anonymisation=scan.summary())))
    with path.open("rb") as stream:
        values = load_workbook(stream, data_only=True)
    provider_comparison = [dict(zip(("provider", "build_jh", "run_monthly_jh", "cloud_monthly_eur", "horizon_years", "tco_eur"),
                                   [values["06_Scenarios"].cell(row, column).value for column in range(1, 7)]))
                           for row in comparison_rows.values()]
    return dict(recalculation=recalculation.to_dict(), validation=report.to_dict(), provider_comparison=provider_comparison,
                anonymisation_matches=scan.match_count, diagram_inventory_delta=0)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=Path(__file__).parent)
    parser.add_argument("--require-native", action="store_true", help="Fail if native Excel/LibreOffice recalculation is unavailable")
    args = parser.parse_args()
    print(json.dumps(generate(args.output_dir, args.require_native), ensure_ascii=False))
