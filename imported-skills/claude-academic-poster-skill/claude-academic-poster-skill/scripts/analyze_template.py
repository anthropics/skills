#!/usr/bin/env python3
"""Analyze a PPTX poster template: dump shape inventory, extract colors, save logos.

Usage:
    python analyze_template.py template.pptx [--outdir analysis/]

Outputs (all in --outdir):
    shapes.txt          Full shape inventory with positions, sizes, types, text, colors
    colors.json         Extracted color palette (RGB hex values per shape)
    logos/              Embedded images extracted from the template
    template_render.png Visual render of the template (requires libreoffice)
"""

import argparse
import json
import os
import re
import subprocess
import sys

from lxml import etree


def analyze(template_path, outdir):
    from pptx import Presentation
    from pptx.util import Inches

    os.makedirs(outdir, exist_ok=True)
    os.makedirs(os.path.join(outdir, "logos"), exist_ok=True)

    prs = Presentation(template_path)
    slide = prs.slides[0]

    slide_w_in = prs.slide_width / 914400
    slide_h_in = prs.slide_height / 914400

    lines = []
    lines.append(f"Template: {template_path}")
    lines.append(f"Slide dimensions: {slide_w_in:.2f}in x {slide_h_in:.2f}in")
    lines.append(f"Total shapes: {len(slide.shapes)}")
    lines.append("")

    all_colors = {}
    design_shapes = []  # shapes that look like design elements (no text, fills)
    content_shapes = []  # shapes with text content

    for i, shape in enumerate(slide.shapes):
        left_in = shape.left / 914400
        top_in = shape.top / 914400
        w_in = shape.width / 914400
        h_in = shape.height / 914400

        line = (
            f"Shape {i:2d} | type={shape.shape_type!s:20s} | "
            f"name=\"{shape.name}\" | "
            f"pos=({left_in:.2f}, {top_in:.2f}) | "
            f"size=({w_in:.2f} x {h_in:.2f})"
        )
        lines.append(line)

        # Extract colors from XML
        xml = etree.tostring(shape._element, pretty_print=True).decode()
        colors = re.findall(r'srgbClr val="([^"]+)"', xml)
        if colors:
            all_colors[f"shape_{i}_{shape.name}"] = list(set(colors))
            lines.append(f"         Colors: {list(set(colors))}")

        # Extract text with formatting
        if hasattr(shape, "text") and shape.text:
            text_preview = shape.text[:150].replace("\n", " | ")
            lines.append(f"         Text: \"{text_preview}\"")

            if hasattr(shape, "text_frame"):
                for pi, para in enumerate(shape.text_frame.paragraphs):
                    for ri, run in enumerate(para.runs):
                        f = run.font
                        sz = f.size / 12700 if f.size else "inherited"
                        try:
                            c = str(f.color.rgb)
                        except Exception:
                            c = "theme"
                        lines.append(
                            f"         Run: size={sz}pt bold={f.bold} "
                            f"color={c} font={f.name} "
                            f"text=\"{run.text[:60]}\""
                        )
            content_shapes.append(i)
        else:
            # No text — likely a design element
            if shape.shape_type in (1, 9):  # AUTO_SHAPE, LINE
                design_shapes.append(i)

        # Extract embedded images
        if shape.shape_type == 13:  # PICTURE
            try:
                img_data = shape.image.blob
                ext = shape.image.content_type.split("/")[-1]
                if ext == "svg+xml":
                    ext = "svg"
                fname = f"logos/shape_{i}_{shape.name.replace(' ', '_')}.{ext}"
                fpath = os.path.join(outdir, fname)
                with open(fpath, "wb") as f:
                    f.write(img_data)
                lines.append(f"         Image saved: {fname} ({len(img_data)} bytes)")
            except Exception as e:
                lines.append(f"         Image: linked/not embedded ({e})")
                design_shapes.append(i)

        lines.append("")

    # Summary
    lines.append("=" * 60)
    lines.append("SUMMARY")
    lines.append("=" * 60)
    lines.append(f"Design elements (keep these): {design_shapes}")
    lines.append(f"Content shapes (replace these): {content_shapes}")
    lines.append("")
    lines.append("Suggested keep_indices for build script:")
    lines.append(f"  keep = {set(design_shapes)}")

    # Write outputs
    shapes_path = os.path.join(outdir, "shapes.txt")
    with open(shapes_path, "w") as f:
        f.write("\n".join(lines))
    print(f"Shape inventory: {shapes_path}")

    colors_path = os.path.join(outdir, "colors.json")
    with open(colors_path, "w") as f:
        json.dump(all_colors, f, indent=2)
    print(f"Color palette: {colors_path}")

    # Render template visually
    try:
        subprocess.run(
            ["libreoffice", "--headless", "--convert-to", "png",
             template_path, "--outdir", outdir],
            capture_output=True, timeout=60,
        )
        # Rename to standard name
        base = os.path.splitext(os.path.basename(template_path))[0]
        src = os.path.join(outdir, f"{base}.png")
        dst = os.path.join(outdir, "template_render.png")
        if os.path.exists(src):
            os.rename(src, dst)
            print(f"Template render: {dst}")
    except Exception as e:
        print(f"Could not render template (libreoffice not available): {e}")

    print(f"\nSlide: {slide_w_in:.2f}in x {slide_h_in:.2f}in")
    print(f"Design shapes to keep: {design_shapes}")
    print(f"Content shapes to replace: {content_shapes}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Analyze a PPTX poster template")
    parser.add_argument("template", help="Path to the .pptx template")
    parser.add_argument("--outdir", default="analysis", help="Output directory")
    args = parser.parse_args()
    analyze(args.template, args.outdir)
