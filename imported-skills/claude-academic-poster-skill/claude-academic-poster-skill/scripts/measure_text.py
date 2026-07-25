#!/usr/bin/env python3
"""Measure actual rendered text heights for poster layout planning.

This script solves the #1 token-burning problem in poster creation:
predicting how much vertical space text will consume. It creates a
temporary PPTX with your actual text at the target column width,
renders it via LibreOffice, and measures the pixel heights.

Usage:
    # Measure specific text items from a JSON file
    python measure_text.py --texts texts.json --width 11.1

    # Quick measure of a single text string
    python measure_text.py --quick "Your bullet text here" --width 11.1 --size 24

The texts.json format:
[
    {"text": "▸ Attention bottleneck limits **long-range** dependency modeling", "size": 24, "bold_markers": true},
    {"text": "▸ Existing methods scale quadratically with sequence length", "size": 24},
    {"group": "problem_bullets", "items": [
        "Self-attention: O(n²) complexity limits scalability",
        "Sparse methods: trade accuracy for efficiency at long contexts",
        "Linear variants: lose fine-grained token interactions"
    ], "size": 24, "line_spacing": 1.35, "bullet": true}
]

Output: measured heights in inches for each item/group, plus a summary JSON.
"""

import argparse
import json
import os
import subprocess
import sys
import tempfile


def create_measurement_pptx(items, column_width_in, template_path=None):
    """Create a PPTX with text items spaced far apart for measurement."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE

    if template_path:
        prs = Presentation(template_path)
        slide = prs.slides[0]
        # Clear all shapes
        for shape in list(slide.shapes):
            shape._element.getparent().remove(shape._element)
    else:
        prs = Presentation()
        prs.slide_width = Inches(24)
        prs.slide_height = Inches(36)
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    cw = Inches(column_width_in)
    spacing = Inches(4.0)  # far apart so they don't overlap
    marker_color = RGBColor(0xFF, 0x00, 0x00)

    for i, item in enumerate(items):
        y = Inches(0.5) + i * spacing

        # Red marker line at the start position
        marker = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                         Inches(0.3), y, cw + Inches(0.4), Inches(0.02))
        marker.fill.solid()
        marker.fill.fore_color.rgb = marker_color
        marker.line.fill.background()

        # Label
        label = slide.shapes.add_textbox(Inches(0.1), y - Inches(0.3),
                                          Inches(3), Inches(0.3))
        tf = label.text_frame
        r = tf.paragraphs[0].add_run()
        r.text = f"[{i}] {item.get('name', item.get('group', 'item'))}"
        r.font.size = Pt(14)
        r.font.color.rgb = marker_color

        if "group" in item:
            # Bullet group — simulate actual poster bullets
            box = slide.shapes.add_textbox(Inches(0.5), y, cw, Inches(3.5))
            tf = box.text_frame
            tf.word_wrap = True
            tf.auto_size = None

            sz = item.get("size", 24)
            sp = item.get("line_spacing", 1.35)
            bullet_char = "▸ " if item.get("bullet", True) else ""

            for j, text in enumerate(item["items"]):
                if j == 0:
                    p = tf.paragraphs[0]
                else:
                    p = tf.add_paragraph()
                    p.space_before = Pt(5)

                p.line_spacing = sp
                p.alignment = PP_ALIGN.LEFT

                if bullet_char:
                    br = p.add_run()
                    br.text = bullet_char
                    br.font.size = Pt(sz)
                    br.font.bold = True
                    br.font.color.rgb = RGBColor(0x61, 0xD8, 0x4E)
                    br.font.name = "Calibri"

                # Handle **bold** markers
                if item.get("bold_markers", True):
                    parts = text.split("**")
                    for k, part in enumerate(parts):
                        r = p.add_run()
                        r.text = part
                        r.font.size = Pt(sz)
                        r.font.bold = (k % 2 == 1)
                        r.font.name = "Calibri"
                else:
                    r = p.add_run()
                    r.text = text
                    r.font.size = Pt(sz)
                    r.font.name = "Calibri"
        else:
            # Single text item
            box = slide.shapes.add_textbox(Inches(0.5), y, cw, Inches(3.0))
            tf = box.text_frame
            tf.word_wrap = True
            tf.auto_size = None

            p = tf.paragraphs[0]
            p.line_spacing = item.get("line_spacing", 1.35)

            sz = item.get("size", 24)
            text = item["text"]

            if item.get("bold_markers", False):
                parts = text.split("**")
                for k, part in enumerate(parts):
                    r = p.add_run()
                    r.text = part
                    r.font.size = Pt(sz)
                    r.font.bold = (k % 2 == 1)
                    r.font.name = "Calibri"
            else:
                r = p.add_run()
                r.text = text
                r.font.size = Pt(sz)
                r.font.bold = item.get("bold", False)
                r.font.name = "Calibri"

    return prs


def measure_from_render(render_path, n_items, spacing_px):
    """Measure text heights from the rendered image by finding red markers."""
    from PIL import Image
    import numpy as np

    img = Image.open(render_path)
    arr = np.array(img)

    # Find red marker lines (R > 200, G < 50, B < 50)
    red_mask = (arr[:, :, 0] > 200) & (arr[:, :, 1] < 50) & (arr[:, :, 2] < 50)

    # Find rows with significant red
    red_rows = np.where(red_mask.sum(axis=1) > 50)[0]

    if len(red_rows) == 0:
        print("WARNING: No red markers found in render")
        return {}

    # Cluster red rows into markers (gaps > 10px = new marker)
    markers = []
    current_start = red_rows[0]
    for j in range(1, len(red_rows)):
        if red_rows[j] - red_rows[j - 1] > 10:
            markers.append(current_start)
            current_start = red_rows[j]
    markers.append(current_start)

    # Height of the full image corresponds to the slide height
    # We know markers are spaced 4.0 inches apart
    # So pixels_per_inch ≈ spacing between first two markers / 4.0
    if len(markers) >= 2:
        px_per_inch = (markers[1] - markers[0]) / 4.0
    else:
        px_per_inch = img.height / 36.0  # fallback

    heights = {}
    for i in range(min(n_items, len(markers))):
        marker_y = markers[i]
        # Scan downward from marker to find where content ends
        # (look for when the row becomes mostly white again)
        content_start = marker_y + 5  # skip the marker line itself

        # Find the bottom of text content (rows that have non-white pixels)
        row_darkness = []
        scan_end = min(marker_y + int(4.0 * px_per_inch), img.height)
        for row in range(content_start, scan_end):
            # Count non-white pixels (any channel < 240)
            dark = np.any(arr[row, :, :3] < 240, axis=1).sum()
            row_darkness.append(dark)

        # Find last row with significant content (> 20 dark pixels)
        last_content = 0
        for j, d in enumerate(row_darkness):
            if d > 20:
                last_content = j

        content_height_px = last_content
        content_height_in = content_height_px / px_per_inch

        heights[i] = round(content_height_in, 2)

    return heights


def main():
    parser = argparse.ArgumentParser(description="Measure text heights for poster layout")
    parser.add_argument("--texts", help="JSON file with text items to measure")
    parser.add_argument("--quick", help="Quick measure a single text string")
    parser.add_argument("--width", type=float, default=11.1,
                        help="Column width in inches (default: 11.1)")
    parser.add_argument("--size", type=int, default=24,
                        help="Font size in pt (for --quick, default: 24)")
    parser.add_argument("--template", help="Template PPTX (for matching slide dimensions)")
    parser.add_argument("--outdir", default="figures/measure", help="Output directory")
    args = parser.parse_args()

    os.makedirs(args.outdir, exist_ok=True)

    if args.quick:
        items = [{"text": args.quick, "size": args.size, "name": "quick"}]
    elif args.texts:
        with open(args.texts) as f:
            items = json.load(f)
    else:
        print("Specify --texts or --quick")
        sys.exit(1)

    # Create measurement PPTX
    prs = create_measurement_pptx(items, args.width, args.template)
    pptx_path = os.path.join(args.outdir, "measure.pptx")
    prs.save(pptx_path)

    # Render via LibreOffice
    subprocess.run(
        ["libreoffice", "--headless", "--convert-to", "pdf",
         pptx_path, "--outdir", args.outdir],
        capture_output=True, timeout=60,
    )
    pdf_path = os.path.join(args.outdir, "measure.pdf")
    if not os.path.exists(pdf_path):
        print("ERROR: LibreOffice rendering failed")
        sys.exit(1)

    # Render PDF to image
    import fitz
    from PIL import Image
    import io

    doc = fitz.open(pdf_path)
    page = doc[0]
    mat = fitz.Matrix(3.0, 3.0)
    pix = page.get_pixmap(matrix=mat)
    img = Image.open(io.BytesIO(pix.tobytes("png")))
    render_path = os.path.join(args.outdir, "measure_render.png")
    img.save(render_path)

    # Measure heights
    heights = measure_from_render(render_path, len(items), spacing_px=0)

    # Print results
    print(f"\nText heights at {args.width}in column width:")
    print("=" * 50)
    results = []
    for i, item in enumerate(items):
        name = item.get("name", item.get("group", f"item_{i}"))
        h = heights.get(i, "?")
        recommended = round(h * 1.15, 1) if isinstance(h, (int, float)) else "?"
        print(f"  [{i}] {name}: {h}in (recommend allocation: {recommended}in)")
        results.append({"index": i, "name": name, "height_in": h, "recommended_in": recommended})

    # Save results
    results_path = os.path.join(args.outdir, "measurements.json")
    with open(results_path, "w") as f:
        json.dump({"column_width_in": args.width, "measurements": results}, f, indent=2)
    print(f"\nResults saved to {results_path}")


if __name__ == "__main__":
    main()
